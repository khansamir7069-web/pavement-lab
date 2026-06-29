import sys
import json
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.db.repository import Database
from app.core import (
    StructuralInput,
    compute_structural_design,
    TrafficInput,
    compute_traffic_analysis,
    PavementLayer,
)
from app.reports.report_builder import CombinedReportContext, build_combined_report
from docx import Document

# Initialize a clean test database
db_path = Path("validation_reports/validation_pavement.db")
if db_path.exists():
    db_path.unlink()

db = Database(db_path)

# Let's define the 3 valid scenarios
scenarios = [
    {
        "id": 1,
        "name": "Low traffic + medium CBR",
        "initial_cvpd": 100.0,
        "growth_rate_pct": 5.0,
        "design_life_years": 10,
        "vdf": 2.0,
        "ldf": 0.75,
        "subgrade_cbr_pct": 5.0,
        "work_name": "Low Traffic + Medium CBR Demo",
        "client": "Client ABC"
    },
    {
        "id": 2,
        "name": "High traffic + low CBR",
        "initial_cvpd": 3000.0,
        "growth_rate_pct": 7.5,
        "design_life_years": 15,
        "vdf": 4.5,
        "ldf": 0.75,
        "subgrade_cbr_pct": 2.5,
        "work_name": "High Traffic + Low CBR Demo",
        "client": "Client DEF"
    },
    {
        "id": 3,
        "name": "Very high MSA case",
        "initial_cvpd": 8000.0,
        "growth_rate_pct": 8.0,
        "design_life_years": 20,
        "vdf": 5.5,
        "ldf": 0.75,
        "subgrade_cbr_pct": 4.0,
        "work_name": "Very High MSA Case Demo",
        "client": "Client GHI"
    }
]

for sc in scenarios:
    print(f"\n--- Processing Scenario {sc['id']}: {sc['name']} ---")
    
    # 1. Create client and project
    c = db.upsert_client(name=sc["client"])
    p = db.create_project(
        work_name=sc["work_name"],
        client_id=c.id,
        agency="Arcstone Infrastructure",
        submitted_by="QA Validator",
        mix_type="DBM-II"
    )
    print(f"Project created with ID: {p.id}")
    
    # 2. Run structural design calculation
    struct_inp = StructuralInput(
        initial_cvpd=sc["initial_cvpd"],
        growth_rate_pct=sc["growth_rate_pct"],
        design_life_years=sc["design_life_years"],
        vdf=sc["vdf"],
        ldf=sc["ldf"],
        subgrade_cbr_pct=sc["subgrade_cbr_pct"]
    )
    struct_res = compute_structural_design(struct_inp)
    db.save_structural_design(project_id=p.id, result=struct_res)
    db.set_module_status(p.id, "structural", "complete")
    print(f"Saved structural design. Calculated MSA: {struct_res.design_msa:.2f}, Subgrade Mr: {struct_res.subgrade_mr_mpa:.2f} MPa")
    print(f"Pavement layers suggest: {', '.join(f'{l.name} ({l.thickness_mm}mm)' for l in struct_res.composition)}")
    
    # 3. Run traffic design calculation
    traffic_inp = TrafficInput(
        initial_cvpd=sc["initial_cvpd"],
        growth_rate_pct=sc["growth_rate_pct"],
        design_life_years=sc["design_life_years"],
        vdf=sc["vdf"],
        ldf=sc["ldf"]
    )
    traffic_res = compute_traffic_analysis(traffic_inp)
    db.save_traffic_analysis(project_id=p.id, result=traffic_res)
    db.set_module_status(p.id, "traffic", "complete")
    print(f"Saved traffic analysis. Calculated MSA: {traffic_res.design_msa:.2f}")

    # 4. Generate combined report Context
    ctx = CombinedReportContext(
        project_title=f"SamPave Report - {sc['name']}",
        work_name=sc["work_name"],
        client=sc["client"],
        agency="Arcstone Infrastructure",
        submitted_by="QA Validator"
    )
    
    # 5. Build combined report
    out_file = Path(f"validation_reports/scenario_{sc['id']}_report.docx")
    out_path, included = build_combined_report(
        out_file,
        db,
        p.id,
        ctx
    )
    print(f"Report generated successfully: {out_file} (Size: {out_file.stat().st_size} bytes)")
    print(f"Included sections: {included}")
    
    # 6. Parse and validate report docx contents
    doc = Document(str(out_file))
    paragraphs = [p_item.text for p_item in doc.paragraphs]
    tables_data = []
    for table in doc.tables:
        for row in table.rows:
            tables_data.append([cell.text.strip() for cell in row.cells])
            
    # Check for presence of key words
    print("Validating content...")
    
    # MSA check
    has_msa_val = False
    for row in tables_data:
        row_str = " | ".join(row)
        if "Cumulative Design Traffic" in row_str or "MSA" in row_str or "Traffic" in row_str:
            msa_formatted = f"{struct_res.design_msa:.2f}"
            if msa_formatted in row_str or f"{struct_res.design_msa:.3f}" in row_str or f"{int(struct_res.design_msa)}" in row_str:
                has_msa_val = True
                
    # Subgrade Mr check
    has_mr_val = False
    for row in tables_data:
        row_str = " | ".join(row)
        if "Subgrade Modulus" in row_str or "Mr" in row_str or "Resilient Modulus" in row_str:
            mr_formatted = f"{struct_res.subgrade_mr_mpa:.2f}"
            if mr_formatted in row_str or f"{int(struct_res.subgrade_mr_mpa)}" in row_str:
                has_mr_val = True

    # Layer thickness checks
    has_layers = True
    for layer in struct_res.composition:
        layer_found = False
        for row in tables_data:
            row_str = " | ".join(row)
            if layer.name in row_str and f"{layer.thickness_mm}" in row_str:
                layer_found = True
        if not layer_found:
            # Let's check paragraphs as well
            for p_text in paragraphs:
                if layer.name in p_text and f"{layer.thickness_mm}" in p_text:
                    layer_found = True
        if not layer_found:
            has_layers = False
            print(f"Warning: layer {layer.name} ({layer.thickness_mm}mm) not found in tables or paragraphs!")
            
    # Check for blank sections / mock certification
    has_blanks = False
    for p_text in paragraphs:
        if "TODO" in p_text:
            has_blanks = True
            
    print(f"Validation summary: MSA verified: {has_msa_val}, Subgrade Mr verified: {has_mr_val}, Layers verified: {has_layers}, Has placeholder/blanks: {has_blanks}")
