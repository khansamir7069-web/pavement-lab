"""Review models for persisted IITPAVE schema diagnostics history.

Phase 34 exposes previously persisted schema diagnostics for operator review.
It is audit-only: rows are recalled as traceability summaries, not engineering
calculation inputs.
"""
from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ._docx_common import (
    add_heading,
    add_kv_table,
    add_note,
    add_p,
    add_signature_block,
    add_table,
    new_portrait_document,
)


IITPAVE_SCHEMA_HISTORY_STATUS_EMPTY = "history_empty"
IITPAVE_SCHEMA_HISTORY_STATUS_AVAILABLE = "history_available_calculations_blocked"
IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_INCLUDED = "schema_history_report_included"
IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_EMPTY = "schema_history_report_empty"


def _payload(raw: Any) -> dict[str, Any]:
    if isinstance(raw, Mapping):
        return dict(raw)
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return {}
    return dict(parsed) if isinstance(parsed, Mapping) else {}


def _dt_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat(sep=" ")
    return str(value or "")


def _diagnostic_rows(summary_payload: Mapping[str, Any]) -> tuple[dict[str, str], ...]:
    report_summary = summary_payload.get("report_summary")
    if not isinstance(report_summary, Mapping):
        return ()
    raw_rows = report_summary.get("diagnostic_rows")
    if not isinstance(raw_rows, list):
        return ()
    rows: list[dict[str, str]] = []
    for item in raw_rows:
        if not isinstance(item, Mapping):
            continue
        rows.append({
            "severity": str(item.get("severity") or ""),
            "field": str(item.get("field") or ""),
            "message": str(item.get("message") or ""),
        })
    return tuple(rows)


@dataclass(frozen=True, slots=True)
class IITPaveSchemaDiagnosticsHistoryItem:
    id: int
    project_id: int
    generated_at: str
    fixture_dir: str
    report_path: str
    workflow_status: str
    manifest_status: str
    parser_audit_ready: bool
    engineering_calculations_allowed: bool
    total_fixture_count: int
    verified_fixture_count: int
    mapped_schema_count: int
    blocked_schema_count: int
    unknown_schema_count: int
    unsupported_schema_count: int
    operator_message: str
    diagnostic_rows: tuple[dict[str, str], ...] = ()

    @property
    def label(self) -> str:
        return (
            f"History #{self.id} | {self.generated_at} | "
            f"{self.workflow_status or 'status_unknown'}"
        )

    @property
    def operator_summary(self) -> tuple[str, ...]:
        lines = [
            f"History #{self.id} generated {self.generated_at}.",
            f"Workflow status: {self.workflow_status or 'status_unknown'}.",
            f"Schema manifest status: {self.manifest_status or 'status_unknown'}.",
            f"Parser audit ready: {self.parser_audit_ready}.",
            (
                "Fixture counts: "
                f"total={self.total_fixture_count}, "
                f"verified={self.verified_fixture_count}, "
                f"mapped={self.mapped_schema_count}, "
                f"blocked={self.blocked_schema_count}, "
                f"unknown={self.unknown_schema_count}, "
                f"unsupported={self.unsupported_schema_count}."
            ),
        ]
        if self.fixture_dir:
            lines.append(f"Fixture folder: {self.fixture_dir}")
        if self.report_path:
            lines.append(f"Report path: {self.report_path}")
        if self.operator_message:
            lines.append(self.operator_message)
        lines.append("Engineering calculations remain blocked.")
        return tuple(lines)

    def as_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "project_id": self.project_id,
            "generated_at": self.generated_at,
            "fixture_dir": self.fixture_dir,
            "report_path": self.report_path,
            "workflow_status": self.workflow_status,
            "manifest_status": self.manifest_status,
            "parser_audit_ready": self.parser_audit_ready,
            "engineering_calculations_allowed": self.engineering_calculations_allowed,
            "total_fixture_count": self.total_fixture_count,
            "verified_fixture_count": self.verified_fixture_count,
            "mapped_schema_count": self.mapped_schema_count,
            "blocked_schema_count": self.blocked_schema_count,
            "unknown_schema_count": self.unknown_schema_count,
            "unsupported_schema_count": self.unsupported_schema_count,
            "operator_message": self.operator_message,
            "operator_summary": list(self.operator_summary),
            "diagnostic_rows": list(self.diagnostic_rows),
        }


@dataclass(frozen=True, slots=True)
class IITPaveSchemaDiagnosticsHistoryReview:
    project_id: int
    status: str
    engineering_calculations_allowed: bool
    items: tuple[IITPaveSchemaDiagnosticsHistoryItem, ...] = ()

    @property
    def ok(self) -> bool:
        return self.status == IITPAVE_SCHEMA_HISTORY_STATUS_AVAILABLE

    @property
    def item_count(self) -> int:
        return len(self.items)

    @property
    def latest_item(self) -> IITPaveSchemaDiagnosticsHistoryItem | None:
        return self.items[0] if self.items else None

    @property
    def operator_summary(self) -> tuple[str, ...]:
        if not self.items:
            return (
                "No persisted IITPAVE schema diagnostics history is available for this project.",
                "Engineering calculations remain blocked.",
            )
        latest = self.latest_item
        return (
            f"{len(self.items)} persisted IITPAVE schema diagnostics record(s) available.",
            f"Latest recalled history id: {latest.id if latest else ''}.",
            "Review is audit-only; engineering calculations remain blocked.",
        )

    def as_dict(self) -> dict[str, Any]:
        return {
            "project_id": self.project_id,
            "ok": self.ok,
            "status": self.status,
            "engineering_calculations_allowed": self.engineering_calculations_allowed,
            "item_count": self.item_count,
            "operator_summary": list(self.operator_summary),
            "items": [item.as_dict() for item in self.items],
        }


@dataclass(frozen=True, slots=True)
class IITPaveSchemaHistoryReportContext:
    project_title: str = ""
    work_name: str = ""
    work_order_no: str = ""
    work_order_date: str = ""
    client: str = ""
    agency: str = ""
    submitted_by: str = ""
    lab_name: str = "Pavement Laboratory"
    report_date: str = ""


@dataclass(frozen=True, slots=True)
class IITPaveSchemaHistoryReportSummary:
    review: IITPaveSchemaDiagnosticsHistoryReview
    status: str
    engineering_calculations_allowed: bool
    included_history_count: int = 0
    diagnostic_row_count: int = 0

    @property
    def ok(self) -> bool:
        return self.status == IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_INCLUDED

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "status": self.status,
            "engineering_calculations_allowed": self.engineering_calculations_allowed,
            "included_history_count": self.included_history_count,
            "diagnostic_row_count": self.diagnostic_row_count,
            "review": self.review.as_dict(),
        }


def build_iitpave_schema_history_item(row: Any) -> IITPaveSchemaDiagnosticsHistoryItem:
    summary_payload = _payload(getattr(row, "summary_json", None))
    return IITPaveSchemaDiagnosticsHistoryItem(
        id=int(getattr(row, "id", 0) or 0),
        project_id=int(getattr(row, "project_id", 0) or 0),
        generated_at=_dt_text(getattr(row, "generated_at", "")),
        fixture_dir=str(getattr(row, "fixture_dir", "") or ""),
        report_path=str(getattr(row, "report_path", "") or ""),
        workflow_status=str(getattr(row, "workflow_status", "") or ""),
        manifest_status=str(getattr(row, "manifest_status", "") or ""),
        parser_audit_ready=bool(getattr(row, "parser_audit_ready", False)),
        engineering_calculations_allowed=bool(
            getattr(row, "engineering_calculations_allowed", False)
        ),
        total_fixture_count=int(getattr(row, "total_fixture_count", 0) or 0),
        verified_fixture_count=int(getattr(row, "verified_fixture_count", 0) or 0),
        mapped_schema_count=int(getattr(row, "mapped_schema_count", 0) or 0),
        blocked_schema_count=int(getattr(row, "blocked_schema_count", 0) or 0),
        unknown_schema_count=int(getattr(row, "unknown_schema_count", 0) or 0),
        unsupported_schema_count=int(getattr(row, "unsupported_schema_count", 0) or 0),
        operator_message=str(getattr(row, "operator_message", "") or ""),
        diagnostic_rows=_diagnostic_rows(summary_payload),
    )


def build_iitpave_schema_history_review(
    project_id: int,
    rows: tuple[Any, ...] | list[Any],
) -> IITPaveSchemaDiagnosticsHistoryReview:
    items = tuple(build_iitpave_schema_history_item(row) for row in rows)
    return IITPaveSchemaDiagnosticsHistoryReview(
        project_id=project_id,
        status=(
            IITPAVE_SCHEMA_HISTORY_STATUS_AVAILABLE
            if items
            else IITPAVE_SCHEMA_HISTORY_STATUS_EMPTY
        ),
        engineering_calculations_allowed=False,
        items=items,
    )


def format_iitpave_schema_history_item_text(
    item: IITPaveSchemaDiagnosticsHistoryItem,
) -> str:
    diagnostics = [
        f"[{row['severity']}] {row['field']}: {row['message']}"
        for row in item.diagnostic_rows
        if row.get("severity") or row.get("field") or row.get("message")
    ]
    blocks = ["\n".join(item.operator_summary)]
    if diagnostics:
        blocks.append("Diagnostics:\n" + "\n".join(diagnostics))
    return "\n\n".join(blocks)


def build_iitpave_schema_history_report_summary(
    review: IITPaveSchemaDiagnosticsHistoryReview,
) -> IITPaveSchemaHistoryReportSummary:
    return IITPaveSchemaHistoryReportSummary(
        review=review,
        status=(
            IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_INCLUDED
            if review.items
            else IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_EMPTY
        ),
        engineering_calculations_allowed=False,
        included_history_count=review.item_count,
        diagnostic_row_count=sum(len(item.diagnostic_rows) for item in review.items),
    )


def _history_rows(
    review: IITPaveSchemaDiagnosticsHistoryReview,
) -> list[list[str]]:
    if not review.items:
        return [["No persisted schema diagnostics history.", "", "", "", "", "No"]]
    return [
        [
            f"#{item.id}",
            item.generated_at,
            item.workflow_status or "status_unknown",
            item.manifest_status or "status_unknown",
            "Yes" if item.parser_audit_ready else "No",
            "No",
        ]
        for item in review.items
    ]


def _fixture_count_rows(item: IITPaveSchemaDiagnosticsHistoryItem) -> list[list[str]]:
    return [
        ["Total fixtures", str(item.total_fixture_count)],
        ["Verified fixtures", str(item.verified_fixture_count)],
        ["Mapped schema fixtures", str(item.mapped_schema_count)],
        ["Blocked schema fixtures", str(item.blocked_schema_count)],
        ["Unknown schema fixtures", str(item.unknown_schema_count)],
        ["Unsupported fixture contracts", str(item.unsupported_schema_count)],
    ]


def _diagnostic_table_rows(item: IITPaveSchemaDiagnosticsHistoryItem) -> list[list[str]]:
    rows = [
        [row["severity"], row["field"], row["message"]]
        for row in item.diagnostic_rows
        if row.get("severity") or row.get("field") or row.get("message")
    ]
    if not rows:
        rows.append(["info", "schema_history", "No blocking diagnostics recorded."])
    return rows


def write_iitpave_schema_history_section(
    doc: Document,
    ctx: IITPaveSchemaHistoryReportContext,
    review: IITPaveSchemaDiagnosticsHistoryReview,
    *,
    include_header: bool = True,
) -> IITPaveSchemaHistoryReportSummary:
    """Append audit-only recalled IITPAVE schema diagnostics history to a report."""
    summary = build_iitpave_schema_history_report_summary(review)

    if include_header:
        add_heading(
            doc,
            "IITPAVE SCHEMA DIAGNOSTICS HISTORY",
            level=1,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        if ctx.project_title:
            add_p(doc, ctx.project_title, bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        date_text = ctx.report_date or datetime.now().strftime("%d-%b-%Y")
        add_p(doc, f"{ctx.lab_name}  -  Report Date: {date_text}",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_heading(doc, "Project Information", level=2)
        add_kv_table(doc, (
            ("Name of Work", ctx.work_name),
            ("Work Order No.", ctx.work_order_no),
            ("Work Order Date", ctx.work_order_date),
            ("Client", ctx.client),
            ("Agency", ctx.agency),
            ("Submitted By", ctx.submitted_by),
        ))

    add_heading(doc, "Audit-Only Recall Summary", level=2)
    add_kv_table(doc, (
        ("Review status", review.status),
        ("Persisted history records", str(review.item_count)),
        ("Engineering calculations allowed", "No"),
    ))
    for line in review.operator_summary:
        add_p(doc, line, size=10)

    add_heading(doc, "Persisted History Records", level=2)
    add_table(
        doc,
        ["History", "Generated", "Workflow status", "Manifest status", "Audit-ready", "Calculations"],
        _history_rows(review),
    )

    if review.latest_item is not None:
        latest = review.latest_item
        add_heading(doc, "Latest Recalled Diagnostic Summary", level=2)
        add_kv_table(doc, (
            ("History ID", f"#{latest.id}"),
            ("Fixture folder", latest.fixture_dir),
            ("Report path", latest.report_path),
            ("Workflow status", latest.workflow_status),
            ("Manifest status", latest.manifest_status),
            ("Parser audit-ready", "Yes" if latest.parser_audit_ready else "No"),
            ("Engineering calculations allowed", "No"),
        ))
        add_table(doc, ["Count", "Value"], _fixture_count_rows(latest))
        if latest.operator_message:
            add_heading(doc, "Original Operator Message", level=3)
            for line in latest.operator_message.splitlines():
                add_p(doc, line, size=9)
        add_heading(doc, "Propagated Diagnostics", level=3)
        add_table(doc, ["Severity", "Field", "Message"], _diagnostic_table_rows(latest))

    add_note(
        doc,
        "This section recalls persisted IITPAVE schema diagnostics for audit "
        "traceability only. It does not extract strains, run mechanistic "
        "calculations, evaluate fatigue or rutting, make IRC compliance "
        "claims, or issue engineering recommendations.",
    )
    return summary


def build_iitpave_schema_history_docx(
    out_path: Path,
    ctx: IITPaveSchemaHistoryReportContext,
    review: IITPaveSchemaDiagnosticsHistoryReview,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = new_portrait_document()
    write_iitpave_schema_history_section(doc, ctx, review, include_header=True)
    add_signature_block(doc)
    doc.save(out_path)
    return out_path
