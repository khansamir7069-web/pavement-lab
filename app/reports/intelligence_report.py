"""Engineering Intelligence Review report section writer."""
from __future__ import annotations

from docx import Document
from app.core.intelligence_checker import IntelligenceResult
from ._docx_common import add_heading, add_p, add_table, add_kv_table


def write_intelligence_section(
    doc: Document,
    intel: IntelligenceResult,
    *,
    section_title: str,
    include_header: bool = True,
) -> None:
    """Write the engineering intelligence section into a docx Document."""
    if include_header:
        add_heading(doc, f"ENGINEERING INTELLIGENCE REVIEW - {section_title}", level=2)
    
    # 1. Health Score and Risk Level
    add_kv_table(doc, (
        ("Engineering Screening Score", f"{intel.health_score:.0f} / 100"),
        ("Risk Classification", intel.risk_level),
    ))
    
    # 2. Compatibility warnings
    add_heading(doc, "Design Compatibility Screening Alerts", level=3)
    if intel.warnings:
        warn_rows = [[f"{idx+1}", w] for idx, w in enumerate(intel.warnings)]
        add_table(doc, ["#", "Engineering compatibility screening alerts"], warn_rows)
    else:
        add_p(doc, "No design compatibility alerts flagged. The layer stack modular ratios and thicknesses are within standard screening ranges.")
        
    # 3. Engineer review required notes
    add_heading(doc, "Engineering Review Remarks", level=3)
    if intel.review_notes:
        for note in intel.review_notes:
            add_p(doc, f"• {note}")
    else:
        add_p(doc, "Standard design checks complete. Pavement structure shows high compliance.")
        
    # 4. Safety disclaimer
    add_p(
        doc,
        "Disclaimer: This Engineering Intelligence Review is a decision-support screening tool. "
        "Final design acceptance requires independent review and sign-off by a qualified pavement engineer. "
        "The Engineering Screening Score is a mathematical index based on typical ranges and does not "
        "constitute final safety approval or certification.",
        italic=True,
        size=9,
    )
