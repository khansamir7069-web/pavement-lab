"""IITPAVE fixture schema diagnostics report integration.

Phase 31 renders Phase-30 schema manifest summaries for operators. It is an
audit/reporting adapter only: no real IITPAVE engineering output is parsed and
no mechanistic calculation path is enabled.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core import (
    IITPAVE_SCHEMA_MANIFEST_STATUS_AUDIT_READY,
    IITPaveFixtureSchemaManifest,
)
from app.core.config_profiles import (
    VALIDATION_ERROR,
    VALIDATION_INFO,
    VALIDATION_WARNING,
)

from ._docx_common import (
    add_heading,
    add_kv_table,
    add_note,
    add_p,
    add_placeholder_banner,
    add_signature_block,
    add_table,
    new_portrait_document,
)


IITPAVE_SCHEMA_REPORT_STATUS_AUDIT_READY = "schema_report_audit_ready"
IITPAVE_SCHEMA_REPORT_STATUS_BLOCKED = "schema_report_blocked"


@dataclass(frozen=True, slots=True)
class IITPaveSchemaReportContext:
    project_title: str = ""
    work_name: str = ""
    work_order_no: str = ""
    work_order_date: str = ""
    client: str = ""
    agency: str = ""
    submitted_by: str = ""
    lab_name: str = "Pavement Laboratory"
    report_date: str = field(default_factory=lambda: datetime.now().strftime("%d-%b-%Y"))


@dataclass(frozen=True, slots=True)
class IITPaveSchemaDiagnosticRow:
    severity: str
    field: str
    message: str

    def as_dict(self) -> dict[str, str]:
        return {
            "severity": self.severity,
            "field": self.field,
            "message": self.message,
        }


@dataclass(frozen=True, slots=True)
class IITPaveSchemaReportSummary:
    manifest: IITPaveFixtureSchemaManifest
    status: str
    parser_readiness: str
    engineering_calculations_allowed: bool
    operator_summary: tuple[str, ...] = ()
    diagnostic_rows: tuple[IITPaveSchemaDiagnosticRow, ...] = ()

    @property
    def ok(self) -> bool:
        return self.status == IITPAVE_SCHEMA_REPORT_STATUS_AUDIT_READY

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "status": self.status,
            "parser_readiness": self.parser_readiness,
            "engineering_calculations_allowed": self.engineering_calculations_allowed,
            "operator_summary": list(self.operator_summary),
            "diagnostic_rows": [r.as_dict() for r in self.diagnostic_rows],
            "manifest": self.manifest.as_dict(),
        }


def _row(severity: str, field: str, message: str) -> IITPaveSchemaDiagnosticRow:
    return IITPaveSchemaDiagnosticRow(
        severity=severity,
        field=field,
        message=message,
    )


def _readiness_label(manifest: IITPaveFixtureSchemaManifest) -> str:
    if manifest.status == IITPAVE_SCHEMA_MANIFEST_STATUS_AUDIT_READY:
        return "Audit-ready fixture schema coverage; engineering calculations blocked."
    return f"Blocked fixture schema coverage: {manifest.status}."


def _operator_summary(manifest: IITPaveFixtureSchemaManifest) -> tuple[str, ...]:
    return tuple(manifest.audit_summary) + (
        "Report integration is diagnostic-only; parser output remains guarded.",
    )


def _diagnostic_rows(
    manifest: IITPaveFixtureSchemaManifest,
) -> tuple[IITPaveSchemaDiagnosticRow, ...]:
    rows: list[IITPaveSchemaDiagnosticRow] = []
    rows.append(_row(
        VALIDATION_INFO,
        "schema_manifest.status",
        f"Schema manifest status: {manifest.status}.",
    ))
    rows.append(_row(
        VALIDATION_INFO if manifest.parser_audit_ready else VALIDATION_WARNING,
        "schema_manifest.parser_audit_ready",
        f"Parser audit readiness: {manifest.parser_audit_ready}.",
    ))
    rows.append(_row(
        VALIDATION_INFO,
        "schema_manifest.engineering_calculations_allowed",
        f"Engineering calculations allowed: {manifest.engineering_calculations_allowed}.",
    ))

    if manifest.blocked_schema_count:
        rows.append(_row(
            VALIDATION_WARNING,
            "schema_manifest.blocked_schema_count",
            f"Blocked verified schema fixtures: {manifest.blocked_schema_count}.",
        ))
    if manifest.unknown_schema_count:
        rows.append(_row(
            VALIDATION_WARNING,
            "schema_manifest.unknown_schema_count",
            f"Unknown verified schema fixtures: {manifest.unknown_schema_count}.",
        ))
    if manifest.unsupported_schema_count:
        rows.append(_row(
            VALIDATION_WARNING,
            "schema_manifest.unsupported_schema_count",
            f"Unsupported fixture contracts: {manifest.unsupported_schema_count}.",
        ))

    for issue in manifest.issues:
        rows.append(_row(issue.severity, issue.field, issue.message))

    # De-duplicate while preserving first-seen order for stable reports.
    seen: set[tuple[str, str, str]] = set()
    out: list[IITPaveSchemaDiagnosticRow] = []
    for item in rows:
        key = (item.severity, item.field, item.message)
        if key in seen:
            continue
        seen.add(key)
        out.append(item)
    return tuple(out)


def build_iitpave_schema_report_summary(
    manifest: IITPaveFixtureSchemaManifest,
) -> IITPaveSchemaReportSummary:
    """Convert a schema manifest into report-ready audit diagnostics."""
    status = (
        IITPAVE_SCHEMA_REPORT_STATUS_AUDIT_READY
        if manifest.parser_audit_ready
        else IITPAVE_SCHEMA_REPORT_STATUS_BLOCKED
    )
    return IITPaveSchemaReportSummary(
        manifest=manifest,
        status=status,
        parser_readiness=_readiness_label(manifest),
        engineering_calculations_allowed=False,
        operator_summary=_operator_summary(manifest),
        diagnostic_rows=_diagnostic_rows(manifest),
    )


def _family_rows(manifest: IITPaveFixtureSchemaManifest) -> list[list[str]]:
    if not manifest.family_coverage:
        return [["No reviewed schema family coverage.", "0", ""]]
    return [
        [
            item.schema_family,
            str(item.fixture_count),
            ", ".join(item.source_filenames),
        ]
        for item in manifest.family_coverage
    ]


def _entry_rows(manifest: IITPaveFixtureSchemaManifest) -> list[list[str]]:
    rows: list[list[str]] = []
    for entry in manifest.entries:
        rows.append([
            entry.source_filename,
            entry.verification_status,
            entry.contract_status,
            entry.schema_status or "not_evaluated",
            entry.schema_family,
            "Yes" if entry.blocked else "No",
        ])
    if not rows:
        rows.append(["No fixture records found.", "", "", "", "", "Yes"])
    return rows


def _diagnostic_table_rows(
    summary: IITPaveSchemaReportSummary,
) -> list[list[str]]:
    rows = [
        [item.severity, item.field, item.message]
        for item in summary.diagnostic_rows
        if item.severity in (VALIDATION_ERROR, VALIDATION_WARNING)
    ]
    if not rows:
        rows.append([VALIDATION_INFO, "schema_manifest", "No blocking diagnostics."])
    return rows


def write_iitpave_schema_manifest_section(
    doc: Document,
    ctx: IITPaveSchemaReportContext,
    manifest: IITPaveFixtureSchemaManifest,
    *,
    include_header: bool = True,
) -> IITPaveSchemaReportSummary:
    """Append an audit-only IITPAVE schema manifest section to a Word doc."""
    summary = build_iitpave_schema_report_summary(manifest)

    if include_header:
        add_heading(
            doc,
            "IITPAVE FIXTURE SCHEMA DIAGNOSTICS",
            level=1,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        if ctx.project_title:
            add_p(doc, ctx.project_title, bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        add_p(doc, f"{ctx.lab_name}  -  Report Date: {ctx.report_date}",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_heading(doc, "Project Information", level=2)
        add_kv_table(doc, (
            ("Name of Work", ctx.work_name),
            ("Work Order No.", ctx.work_order_no),
            ("Work Order Date", ctx.work_order_date),
            ("Client", ctx.client),
            ("Agency", ctx.agency),
            ("Submitted By", ctx.submitted_by),
        ))

    if not summary.ok:
        add_placeholder_banner(doc, summary.parser_readiness)

    add_heading(doc, "Parser Readiness Summary", level=2)
    add_kv_table(doc, (
        ("Manifest status", manifest.status),
        ("Parser audit-ready", "Yes" if manifest.parser_audit_ready else "No"),
        ("Engineering calculations allowed", "No"),
        ("Verified fixtures", str(manifest.verified_fixture_count)),
        ("Mapped schema fixtures", str(manifest.mapped_schema_count)),
        ("Blocked schema fixtures", str(manifest.blocked_schema_count)),
        ("Unknown schema fixtures", str(manifest.unknown_schema_count)),
        ("Unsupported fixture contracts", str(manifest.unsupported_schema_count)),
    ))

    add_heading(doc, "Operator Audit Summary", level=2)
    for line in summary.operator_summary:
        add_p(doc, line, size=10)

    add_heading(doc, "Schema Family Coverage", level=2)
    add_table(doc, ["Schema family", "Fixtures", "Source files"], _family_rows(manifest))

    add_heading(doc, "Fixture Schema Status", level=2)
    add_table(
        doc,
        ["Fixture", "Verification", "Contract", "Schema status", "Schema family", "Blocked"],
        _entry_rows(manifest),
    )

    add_heading(doc, "Blocking Diagnostics", level=2)
    add_table(doc, ["Severity", "Field", "Message"], _diagnostic_table_rows(summary))
    add_note(
        doc,
        "This section is an audit diagnostic for reviewed IITPAVE fixture schema "
        "coverage only. It does not authorize parser output, mechanistic checks, "
        "or compliance conclusions.",
    )
    return summary


def build_iitpave_schema_manifest_docx(
    out_path: Path,
    ctx: IITPaveSchemaReportContext,
    manifest: IITPaveFixtureSchemaManifest,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = new_portrait_document()
    write_iitpave_schema_manifest_section(doc, ctx, manifest, include_header=True)
    add_signature_block(doc)
    doc.save(out_path)
    return out_path
