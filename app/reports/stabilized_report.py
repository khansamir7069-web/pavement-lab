"""Stabilized Pavement (CTB/CTS) design report section and document builder."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.stabilized_design import StabilizedResult

from ._docx_common import (
    add_heading,
    add_kv_table,
    add_note,
    add_p,
    add_signature_block,
    add_table,
    new_portrait_document,
)


@dataclass(frozen=True, slots=True)
class StabilizedReportContext:
    project_title: str = ""
    work_name: str = ""
    work_order_no: str = ""
    work_order_date: str = ""
    client: str = ""
    agency: str = ""
    submitted_by: str = ""
    lab_name: str = "Pavement Laboratory"
    report_date: str = field(default_factory=lambda: datetime.now().strftime("%d-%b-%Y"))


def write_stabilized_section(
    doc: Document,
    ctx: StabilizedReportContext,
    result: StabilizedResult,
    *,
    include_header: bool = True,
) -> None:
    """Write the stabilized pavement design details into a docx Document."""
    if include_header:
        add_heading(doc, "STABILIZED PAVEMENT DESIGN (CTB/CTS)",
                    level=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        if ctx.project_title:
            add_p(doc, ctx.project_title, bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        add_p(doc, "Design basis: IRC:37-2018 (empirically derived stabilized pavement guidelines).",
              italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_p(doc, f"{ctx.lab_name}  •  Report Date: {ctx.report_date}",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

        add_heading(doc, "Project Information", level=2)
        add_kv_table(doc, (
            ("Name of Work",    ctx.work_name),
            ("Work Order No.",  ctx.work_order_no),
            ("Work Order Date", ctx.work_order_date),
            ("Client",          ctx.client),
            ("Agency",          ctx.agency),
            ("Submitted By",    ctx.submitted_by),
        ))

    inp = result.inputs

    # 1. Validation & Mode Stamp
    add_heading(doc, "1. Design Validation Mode", level=2)
    add_kv_table(doc, (
        ("Validation Mode", f"★ {result.validation_mode} ★"),
        ("CTB Fatigue Status", "Mechanistic verification required" if result.validation_mode == "Decision Support Mode" else "Mechanistically Verified (IITPAVE)"),
    ))
    add_p(
        doc,
        f"IMPORTANT NOTE: {result.safety_disclaimer}",
        bold=True,
        size=10,
    )

    # 2. Design Inputs
    add_heading(doc, "2. Material Inputs & Design Parameters", level=2)
    
    ctb_ucs_display = f"{inp.ctb_ucs_mpa:.1f} MPa" if inp.ctb_ucs_mpa > 0 else "Not specified (Screening alert)"

    add_kv_table(doc, (
        ("CTB Layer Thickness", f"{inp.ctb_thickness_mm:.0f} mm"),
        ("CTB Elastic Modulus (E_CTB)", f"{inp.ctb_modulus_mpa:.0f} MPa"),
        ("CTB 28-day UCS Strength", ctb_ucs_display),
        ("CTB Poisson's Ratio", f"{inp.ctb_poisson:.2f}"),
        ("CTS Strength Class / Grade", inp.cts_class),
        ("CTS Layer Thickness", f"{inp.cts_thickness_mm:.0f} mm"),
        ("CTS Elastic Modulus (E_CTS)", f"{inp.cts_modulus_mpa:.0f} MPa"),
        ("CTS Poisson's Ratio", f"{inp.cts_poisson:.2f}"),
        ("Bituminous Cover Thickness", f"{inp.bituminous_thickness_mm:.0f} mm"),
        ("Granular Sub-base Thickness", f"{inp.gsb_thickness_mm:.0f} mm"),
        ("Baseline Traffic (MSA)", f"{inp.flexible_design_msa:g} MSA"),
        ("Baseline Subgrade CBR (%)", f"{inp.flexible_subgrade_cbr:g} %"),
    ))

    # 3. Pavement Comparison
    add_heading(doc, f"3. {result.comparison_label}", level=2)
    add_p(doc, "Side-by-side comparison of conventional flexible pavement catalogue composition vs. cement-stabilized pavement structure:")

    # Build layers side-by-side table rows
    comparison_rows = []
    max_layers = max(len(result.conventional_composition), len(result.stabilized_composition))
    
    for i in range(max_layers):
        flex_txt = "—"
        stab_txt = "—"
        if i < len(result.conventional_composition):
            ly = result.conventional_composition[i]
            flex_txt = f"{ly.name} ({ly.material}): {ly.thickness_mm:.0f} mm"
        if i < len(result.stabilized_composition):
            ly = result.stabilized_composition[i]
            stab_txt = f"{ly.name} ({ly.material}): {ly.thickness_mm:.0f} mm"
        comparison_rows.append([flex_txt, stab_txt])

    # Totals row
    total_flex = sum(ly.thickness_mm for ly in result.conventional_composition)
    total_stab = sum(ly.thickness_mm for ly in result.stabilized_composition)
    comparison_rows.append([
        f"TOTAL DEPTH: {total_flex:.0f} mm",
        f"TOTAL DEPTH: {total_stab:.0f} mm"
    ])

    add_table(doc, ["Conventional Flexible Design (Catalogue)", "Stabilized Design (CTB/CTS)"], comparison_rows)

    # Thickness savings
    savings_txt = (
        f"Thickness Savings: {result.thickness_savings_mm:.0f} mm ({result.thickness_savings_pct:.1f}% reduction). "
        "Note: Thickness savings are indicative comparison only, not an optimization proof."
    )
    add_p(doc, savings_txt, bold=True, size=11)

    # 4. Engineering Warnings & Screening
    add_heading(doc, "4. Engineering Warnings & Screening Remarks", level=2)
    if result.warnings:
        warn_rows = [[f"{idx+1}", w] for idx, w in enumerate(result.warnings)]
        add_table(doc, ["#", "Engineering screening alerts & placeholders"], warn_rows)
    else:
        add_p(doc, "No engineering screening warnings generated.")

    add_note(
        doc,
        "Disclaimer: This module is intended for structural comparison and decision-support "
        "only. Empirically suggested configurations must be verified using local material "
        "characterization, fatigue trials, and mechanistic strain validation (IITPAVE)."
    )


def build_stabilized_docx(
    out_path: Path,
    ctx: StabilizedReportContext,
    result: StabilizedResult,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = new_portrait_document()
    write_stabilized_section(doc, ctx, result, include_header=True)
    add_signature_block(doc)
    doc.save(out_path)
    return out_path
