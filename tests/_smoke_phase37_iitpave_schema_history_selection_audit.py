"""Phase-37 smoke - report-time IITPAVE schema-history selection audit trail."""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_tmp = Path(tempfile.mkdtemp(prefix="phase37_iitpave_schema_history_selection_audit_"))
_db_path = _tmp / "phase37.db"

import app.config as _cfg  # noqa: E402
_cfg.DB_PATH = _db_path

from docx import Document  # noqa: E402

from app.core import DEFAULT_FIXTURE_MANIFEST, FIXTURE_STATUS_VERIFIED  # noqa: E402
from app.db.project_exchange import export_project, import_project  # noqa: E402
from app.db.repository import Database  # noqa: E402
from app.reports import (  # noqa: E402
    CombinedReportContext,
    IITPAVE_SCHEMA_HISTORY_SELECTION_AUDIT_STATUS_RECORDED,
    IITPAVE_SCHEMA_HISTORY_SELECTION_STATUS_PARTIAL_UNKNOWN,
    build_combined_report,
    build_iitpave_schema_history_report_summary,
    build_iitpave_schema_history_review,
    run_iitpave_schema_diagnostics_workflow,
)


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _fixture_folder(name: str) -> Path:
    folder = _tmp / name
    folder.mkdir(exist_ok=True)
    sample = _write_fixture(folder, f"{name}.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(folder, [{
        "filename": sample.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    return folder


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _persist_history_pair(db: Database, project_id: int) -> tuple:
    first = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("first_fixture"),
        report_path=_tmp / "first_schema_diagnostics.docx",
    )
    second = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("second_fixture"),
        report_path=_tmp / "second_schema_diagnostics.docx",
    )
    row1 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=first)
    row2 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=second)
    return row1, row2


def _safe_text(value) -> str:
    return json.dumps(value, sort_keys=True).lower()


def main() -> int:
    db = Database(_db_path)
    project = db.create_project(
        work_name="Phase 37 History Selection Audit",
        agency="State PWD",
        submitted_by="QA Engineer",
        mix_type="DBM-II",
    )
    row1, row2 = _persist_history_pair(db, project.id)
    rows = db.list_iitpave_schema_diagnostics(project.id)

    print("=== 1) Typed report summary carries audit-only selection trail ===")
    review = build_iitpave_schema_history_review(
        project.id,
        rows,
        selected_history_ids=(999999, row1.id),
    )
    summary = build_iitpave_schema_history_report_summary(
        review,
        report_path=_tmp / "typed_audit.docx",
    )
    assert summary.selection is not None
    assert summary.selection.status == IITPAVE_SCHEMA_HISTORY_SELECTION_STATUS_PARTIAL_UNKNOWN
    assert summary.selection.available_history_ids == (row2.id, row1.id)
    assert summary.selection.selected_history_ids == (row1.id,)
    assert summary.selection_audit_trail is not None
    assert summary.selection_audit_trail.decision_status == summary.selection.status
    assert summary.selection_audit_trail.available_history_ids == (row2.id, row1.id)
    assert summary.selection_audit_trail.selected_history_ids == (row1.id,)
    assert summary.selection_audit_trail.skipped_unknown_history_ids == (999999,)
    assert summary.selection_audit_trail.engineering_calculations_allowed is False
    audit_text = _safe_text(summary.as_dict())
    assert "selection_audit_trail" in audit_text
    assert "report-time iitpave schema history inclusion decision recorded" in audit_text
    assert "fatigue_life" not in audit_text
    assert "rutting_life" not in audit_text
    assert "irc_compliance" not in audit_text
    assert "point_results" not in audit_text
    print("  [PASS] typed audit trail is deterministic and calculation-free")

    print("\n=== 2) Combined report persists the report-time selection decision ===")
    selected_doc = _tmp / "selected_schema_history_audit.docx"
    _out, included = build_combined_report(
        selected_doc,
        db,
        project.id,
        CombinedReportContext(
            project_title=project.work_name,
            work_name=project.work_name,
            agency=project.agency or "",
            submitted_by=project.submitted_by or "",
            mix_type_key=project.mix_type or "",
        ),
        schema_history_selection_ids=(row1.id,),
    )
    assert "IITPAVE Schema Diagnostics History" in included
    audit_rows = db.list_iitpave_schema_history_selection_audits(project.id)
    assert len(audit_rows) == 1
    audit_row = audit_rows[0]
    assert audit_row.report_path == str(selected_doc)
    assert audit_row.decision_status == "operator_history_selected"
    assert json.loads(audit_row.available_history_ids_json) == [row2.id, row1.id]
    assert json.loads(audit_row.selected_history_ids_json) == [row1.id]
    assert json.loads(audit_row.skipped_unknown_history_ids_json) == []
    assert audit_row.included_history_count == 1
    assert audit_row.engineering_calculations_allowed is False
    persisted_text = _safe_text(json.loads(audit_row.summary_json))
    assert "engineering calculations remain blocked" in persisted_text
    assert "fatigue_life" not in persisted_text
    assert "rutting_life" not in persisted_text
    assert "irc_compliance" not in persisted_text
    assert "point_results" not in persisted_text
    print("  [PASS] persisted audit row records the report-time choice safely")

    print("\n=== 3) Report output contains operator-readable audit trail metadata ===")
    doc_text = _doc_text(selected_doc)
    assert "Report-Time Selection Audit Trail" in doc_text
    assert IITPAVE_SCHEMA_HISTORY_SELECTION_AUDIT_STATUS_RECORDED in doc_text
    assert "Selected history IDs" in doc_text
    assert f"#{row1.id}" in doc_text
    assert f"#{row2.id}" not in doc_text
    assert "Engineering calculations remain blocked" in doc_text
    assert "fatigue_life" not in doc_text.lower()
    assert "rutting_life" not in doc_text.lower()
    assert "irc_compliance" not in doc_text.lower()
    assert "point_results" not in doc_text.lower()
    print("  [PASS] report metadata is operator-readable and audit-only")

    print("\n=== 4) Export/import preserves selection-audit serialization ===")
    exported = export_project(db, project.id)
    audit_payloads = exported["records"]["iitpave_schema_history_selection_audits"]
    assert len(audit_payloads) == 1
    assert audit_payloads[0]["decision_status"] == "operator_history_selected"
    imported = import_project(db, exported)
    imported_audits = db.list_iitpave_schema_history_selection_audits(imported.project_id)
    assert len(imported_audits) == 1
    assert imported_audits[0].decision_status == "operator_history_selected"
    imported_summary = _safe_text(json.loads(imported_audits[0].summary_json))
    assert "selection_audit_trail" in imported_summary
    assert "engineering_calculations_allowed" in imported_summary
    assert "fatigue_life" not in imported_summary
    assert "rutting_life" not in imported_summary
    assert "irc_compliance" not in imported_summary
    print("  [PASS] project exchange preserves audit-trail records compatibly")

    print("\nPHASE 37 IITPAVE SCHEMA HISTORY SELECTION AUDIT SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
