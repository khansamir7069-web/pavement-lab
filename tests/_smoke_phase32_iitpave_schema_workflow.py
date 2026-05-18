"""Phase-32 smoke - IITPAVE fixture-folder schema diagnostics workflow."""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_tmp = Path(tempfile.mkdtemp(prefix="phase32_iitpave_schema_workflow_"))
_db_path = _tmp / "phase32.db"

import app.config as _cfg  # noqa: E402
_cfg.DB_PATH = _db_path

from docx import Document  # noqa: E402
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox  # noqa: E402

from app.core import DEFAULT_FIXTURE_MANIFEST, FIXTURE_STATUS_VERIFIED  # noqa: E402
from app.reports import (  # noqa: E402
    IITPAVE_SCHEMA_WORKFLOW_STATUS_REPORT_BLOCKED,
    IITPAVE_SCHEMA_WORKFLOW_STATUS_REPORT_WRITTEN,
    IITPAVE_SCHEMA_WORKFLOW_STATUS_SUMMARY_ONLY,
    IITPaveSchemaReportContext,
    run_iitpave_schema_diagnostics_workflow,
)

import app.db.repository as _repo  # noqa: E402
from app.db.repository import Database  # noqa: E402


_messages: list[tuple[str, str, str]] = []
QMessageBox.information = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("info", title, message)) or 0
)
QMessageBox.warning = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("warning", title, message)) or 0
)
QMessageBox.critical = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("critical", title, message)) or 0
)


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _fixture_folder() -> Path:
    folder = _tmp / "fixtures"
    folder.mkdir(exist_ok=True)
    sample = _write_fixture(folder, "direct_stress_strain.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(folder, [{
        "filename": sample.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    return folder


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main() -> int:
    fixture_dir = _fixture_folder()

    print("=== 1) Workflow can generate summary without report file ===")
    summary_only = run_iitpave_schema_diagnostics_workflow(fixture_dir)
    assert summary_only.status == IITPAVE_SCHEMA_WORKFLOW_STATUS_SUMMARY_ONLY
    assert summary_only.report_written is False
    assert summary_only.engineering_calculations_allowed is False
    assert "Engineering calculations remain blocked" in summary_only.operator_message
    print("  [PASS] summary-only workflow is audit-safe")

    print("\n=== 2) Workflow writes schema diagnostics Word report ===")
    report_path = _tmp / "schema_diagnostics.docx"
    written = run_iitpave_schema_diagnostics_workflow(
        fixture_dir,
        report_path=report_path,
        context=IITPaveSchemaReportContext(project_title="Phase 32 Workflow"),
    )
    assert written.status == IITPAVE_SCHEMA_WORKFLOW_STATUS_REPORT_WRITTEN
    assert written.report_written is True
    assert written.engineering_calculations_allowed is False
    assert report_path.is_file()
    text = _doc_text(report_path)
    assert "IITPAVE FIXTURE SCHEMA DIAGNOSTICS" in text
    assert "Engineering calculations allowed" in text
    print("  [PASS] report workflow emits operator-readable diagnostics")

    print("\n=== 3) Invalid report path is blocked without calculations ===")
    blocked = run_iitpave_schema_diagnostics_workflow(
        fixture_dir,
        report_path=_tmp / "schema_diagnostics.txt",
    )
    assert blocked.status == IITPAVE_SCHEMA_WORKFLOW_STATUS_REPORT_BLOCKED
    assert blocked.report_written is False
    assert blocked.engineering_calculations_allowed is False
    assert "must end with .docx" in blocked.operator_message
    print("  [PASS] invalid report destinations are refused safely")

    print("\n=== 4) MainWindow exposes local fixture-folder workflow ===")
    _repo._singleton = Database(_db_path)
    app = QApplication.instance() or QApplication([])
    from app.ui.main_window import MainWindow  # noqa: E402

    w = MainWindow()
    assert w.btn_iitpave_schema is not None
    ui_report = _tmp / "ui_schema_diagnostics.docx"
    QFileDialog.getExistingDirectory = staticmethod(lambda *a, **k: str(fixture_dir))
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(ui_report), "Word"))
    w.btn_iitpave_schema.click()
    assert ui_report.is_file()
    assert any("IITPAVE schema diagnostics" in item[1] for item in _messages)
    assert "Engineering calculations remain blocked" in _messages[-1][2]
    assert "IITPAVE FIXTURE SCHEMA DIAGNOSTICS" in _doc_text(ui_report)
    w.close()
    app.processEvents()
    print("  [PASS] UI button selects folder and writes guarded diagnostics report")

    print("\n=== 5) Serialization remains audit-only ===")
    payload = written.as_dict()
    serialized = json.dumps(payload).lower()
    assert "fixture_dir" in serialized
    assert "report_summary" in serialized
    assert "engineering_calculations_allowed" in serialized
    assert "fatigue" not in serialized
    assert "rutting" not in serialized
    assert "irc_compliance" not in serialized
    assert "point_results" not in serialized
    assert payload["engineering_calculations_allowed"] is False
    print("  [PASS] workflow serialization carries no engineering results")

    print("\nPHASE 32 IITPAVE SCHEMA WORKFLOW SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
