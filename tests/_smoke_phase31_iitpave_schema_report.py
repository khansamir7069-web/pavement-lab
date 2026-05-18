"""Phase-31 smoke - IITPAVE schema manifest report diagnostics integration."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from docx import Document

from app.core import (
    DEFAULT_FIXTURE_MANIFEST,
    FIXTURE_STATUS_OPERATOR_PENDING,
    FIXTURE_STATUS_VERIFIED,
    build_iitpave_fixture_schema_manifest,
)
from app.reports import (
    IITPAVE_SCHEMA_REPORT_STATUS_AUDIT_READY,
    IITPAVE_SCHEMA_REPORT_STATUS_BLOCKED,
    IITPaveSchemaReportContext,
    build_iitpave_schema_manifest_docx,
    build_iitpave_schema_report_summary,
)


_tmp = Path(tempfile.mkdtemp(prefix="phase31_iitpave_schema_report_"))


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main() -> int:
    print("=== 1) Audit-ready manifest becomes report-ready diagnostics ===")
    ready_dir = _tmp / "ready"
    ready_dir.mkdir()
    ready_fixture = _write_fixture(ready_dir, "direct_stress_strain.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(ready_dir, [{
        "filename": ready_fixture.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    ready_manifest = build_iitpave_fixture_schema_manifest(ready_dir)
    ready_summary = build_iitpave_schema_report_summary(ready_manifest)
    assert ready_summary.status == IITPAVE_SCHEMA_REPORT_STATUS_AUDIT_READY
    assert ready_summary.ok is True
    assert ready_summary.engineering_calculations_allowed is False
    assert "engineering calculations blocked" in ready_summary.parser_readiness
    assert ready_summary.operator_summary
    print("  [PASS] ready manifest is exposed as audit-only report summary")

    print("\n=== 2) Blocked/unsupported schema diagnostics propagate ===")
    blocked_dir = _tmp / "blocked"
    blocked_dir.mkdir()
    unknown = _write_fixture(blocked_dir, "unknown_schema.out", [
        "IITPAVE OUTPUT",
        "Elastic Layer Profile",
        "Layer Depth Radius",
        "1 50 0",
    ])
    unsupported = _write_fixture(blocked_dir, "unsupported.out", [
        "random console log",
        "abc def",
    ])
    pending = _write_fixture(blocked_dir, "pending.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "1 2 3",
    ])
    _write_manifest(blocked_dir, [
        {"filename": unknown.name, "verification_status": FIXTURE_STATUS_VERIFIED},
        {"filename": unsupported.name, "verification_status": FIXTURE_STATUS_VERIFIED},
        {"filename": pending.name, "verification_status": FIXTURE_STATUS_OPERATOR_PENDING},
    ])
    blocked_manifest = build_iitpave_fixture_schema_manifest(blocked_dir)
    blocked_summary = build_iitpave_schema_report_summary(blocked_manifest)
    assert blocked_summary.status == IITPAVE_SCHEMA_REPORT_STATUS_BLOCKED
    assert blocked_summary.ok is False
    assert blocked_summary.engineering_calculations_allowed is False
    row_payload = json.dumps([r.as_dict() for r in blocked_summary.diagnostic_rows]).lower()
    assert "unknown verified schema fixtures" in row_payload
    assert "unsupported fixture contracts" in row_payload
    assert "parser audit readiness: false" in row_payload
    print("  [PASS] blocked and unsupported diagnostics are report-visible")

    print("\n=== 3) Word report section renders operator-readable diagnostics ===")
    out_path = _tmp / "iitpave_schema_diagnostics.docx"
    build_iitpave_schema_manifest_docx(
        out_path,
        IITPaveSchemaReportContext(
            project_title="Phase 31 Schema Diagnostics",
            work_name="Flexible Pavement IITPAVE Audit",
            client="Smoke Client",
        ),
        blocked_manifest,
    )
    assert out_path.is_file()
    text = _doc_text(out_path)
    assert "IITPAVE FIXTURE SCHEMA DIAGNOSTICS" in text
    assert "Parser Readiness Summary" in text
    assert "Operator Audit Summary" in text
    assert "Blocking Diagnostics" in text
    assert "Engineering calculations allowed" in text
    assert "No" in text
    assert "unknown_schema.out" in text
    assert "unsupported.out" in text
    print("  [PASS] Word diagnostics section is generated without parser execution")

    print("\n=== 4) Serialization remains audit-only ===")
    payload = ready_summary.as_dict()
    blocked_payload = blocked_summary.as_dict()
    serialized = json.dumps({"ready": payload, "blocked": blocked_payload}).lower()
    assert "operator_summary" in serialized
    assert "diagnostic_rows" in serialized
    assert "engineering_calculations_allowed" in serialized
    assert "point_results" not in serialized
    assert "strain_value" not in serialized
    assert "epsilon" not in serialized
    assert payload["engineering_calculations_allowed"] is False
    assert blocked_payload["engineering_calculations_allowed"] is False
    print("  [PASS] report serialization carries diagnostics without engineering values")

    print("\nPHASE 31 IITPAVE SCHEMA REPORT SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
