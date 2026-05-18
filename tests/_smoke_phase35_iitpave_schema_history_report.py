"""Phase-35 smoke - report inclusion for IITPAVE schema history."""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

_tmp = Path(tempfile.mkdtemp(prefix="phase35_iitpave_schema_history_report_"))
_db_path = _tmp / "phase35.db"

import app.config as _cfg  # noqa: E402
_cfg.DB_PATH = _db_path

from docx import Document  # noqa: E402

from app.core import DEFAULT_FIXTURE_MANIFEST, FIXTURE_STATUS_VERIFIED  # noqa: E402
from app.db.project_exchange import export_project, import_project  # noqa: E402
from app.db.repository import Database  # noqa: E402
from app.reports import (  # noqa: E402
    CombinedReportContext,
    IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_INCLUDED,
    IITPaveSchemaHistoryReportContext,
    build_combined_report,
    build_iitpave_schema_history_docx,
    build_iitpave_schema_history_report_summary,
    build_iitpave_schema_history_review,
    run_iitpave_schema_diagnostics_workflow,
)


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _fixture_folder() -> Path:
    folder = _tmp / "fixtures"
    folder.mkdir(exist_ok=True)
    sample = _write_fixture(folder, "direct_stress_strain.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(folder, [{
        "filename": sample.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    return folder


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _safe_text(value) -> str:
    return json.dumps(value, sort_keys=True).lower()


def main() -> int:
    fixture_dir = _fixture_folder()
    db = Database(_db_path)
    project = db.create_project(
        work_name="Phase 35 History Report",
        agency="State PWD",
        submitted_by="QA Engineer",
        mix_type="DBM-II",
    )
    result = run_iitpave_schema_diagnostics_workflow(
        fixture_dir,
        report_path=_tmp / "schema_diagnostics.docx",
    )
    db.save_iitpave_schema_diagnostics(project_id=project.id, result=result)
    review = build_iitpave_schema_history_review(
        project.id,
        db.list_iitpave_schema_diagnostics(project.id),
    )

    print("=== 1) Recalled history builds typed report summary safely ===")
    summary = build_iitpave_schema_history_report_summary(review)
    assert summary.status == IITPAVE_SCHEMA_HISTORY_REPORT_STATUS_INCLUDED
    assert summary.ok is True
    assert summary.engineering_calculations_allowed is False
    assert summary.included_history_count == 1
    serialized = _safe_text(summary.as_dict())
    assert '"engineering_calculations_allowed": false' in serialized
    assert "diagnostic_row_count" in serialized
    assert "fatigue" not in serialized
    assert "rutting" not in serialized
    assert "irc_compliance" not in serialized
    assert "point_results" not in serialized
    print("  [PASS] typed report summary remains audit-only")

    print("\n=== 2) Standalone Word section is operator-readable ===")
    standalone = _tmp / "schema_history_section.docx"
    build_iitpave_schema_history_docx(
        standalone,
        IITPaveSchemaHistoryReportContext(
            project_title=project.work_name,
            work_name=project.work_name,
            agency=project.agency or "",
            submitted_by=project.submitted_by or "",
        ),
        review,
    )
    text = _doc_text(standalone)
    assert "IITPAVE SCHEMA DIAGNOSTICS HISTORY" in text
    assert "Audit-Only Recall Summary" in text
    assert "Engineering calculations allowed" in text
    assert "No" in text
    assert "Latest Recalled Diagnostic Summary" in text
    assert "This section recalls persisted IITPAVE schema diagnostics" in text
    print("  [PASS] history report section renders recalled diagnostics")

    print("\n=== 3) Combined report includes recalled history as audit section ===")
    combined = _tmp / "combined_schema_history.docx"
    out, included = build_combined_report(
        combined,
        db,
        project.id,
        CombinedReportContext(
            project_title=project.work_name,
            work_name=project.work_name,
            agency=project.agency or "",
            submitted_by=project.submitted_by or "",
            mix_type_key=project.mix_type or "",
        ),
    )
    assert out == combined
    assert included == ["IITPAVE Schema Diagnostics History"]
    combined_text = _doc_text(combined)
    assert "COMBINED PAVEMENT-DESIGN REPORT" in combined_text
    assert "IITPAVE Schema Diagnostics History" in combined_text
    assert "IITPAVE SCHEMA DIAGNOSTICS HISTORY" in combined_text
    assert "calculations blocked" in combined_text
    assert "fatigue_life" not in combined_text.lower()
    assert "rutting_life" not in combined_text.lower()
    assert "irc_compliance" not in combined_text.lower()
    assert "point_results" not in combined_text.lower()
    print("  [PASS] combined consultancy report includes audit-only history")

    print("\n=== 4) Export/import compatibility preserves report inclusion ===")
    exported = export_project(db, project.id)
    imported = import_project(db, exported)
    imported_project = db.get_project(imported.project_id)
    imported_combined = _tmp / "imported_combined_schema_history.docx"
    _out, imported_included = build_combined_report(
        imported_combined,
        db,
        imported.project_id,
        CombinedReportContext(
            project_title=imported.work_name,
            work_name=imported.work_name,
            agency=imported_project.agency if imported_project else "",
            submitted_by=imported_project.submitted_by if imported_project else "",
            mix_type_key=imported_project.mix_type if imported_project else "",
        ),
    )
    assert imported_included == ["IITPAVE Schema Diagnostics History"]
    assert "IITPAVE SCHEMA DIAGNOSTICS HISTORY" in _doc_text(imported_combined)
    print("  [PASS] imported history remains reportable")

    print("\nPHASE 35 IITPAVE SCHEMA HISTORY REPORT SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
