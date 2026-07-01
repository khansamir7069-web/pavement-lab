from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
import pytest
from app.config import USER_DATA_DIR
from app.core import (
    IITPAVE_RUNNER_EXTERNAL,
    IITPAVE_RUNNER_STUB,
    IITPaveRunnerConfig,
    StabilizedInput,
    compute_stabilized_design,
    run_stabilized_iitpave_mechanistic_workflow,
    load_persisted_config,
    save_persisted_config,
    detect_iitpave_version,
    get_last_run_info,
    validate_iitpave_environment,
    IITPaveExternalExeRunner,
)


IITPAVE_OUTPUT_TABLE = """IITPAVE OUTPUT
Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR
200 0 0.35 0.11 0.09 0 0 -0.000050 0.000120 0.000110
680L 0 0.04 0.01 0.01 0 0 -0.000220 -0.000030 -0.000020
"""


def test_installation_manager_config_persistence() -> None:
    # Test saving and loading config
    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_STUB,
        configured_executable_path="dummy_path.exe",
        include_path_search=False,
        timeout_sec=42.0,
    )
    save_persisted_config(cfg)
    loaded = load_persisted_config()
    assert loaded.mode == IITPAVE_RUNNER_STUB
    assert loaded.configured_executable_path == "dummy_path.exe"
    assert loaded.include_path_search is False
    assert loaded.timeout_sec == 42.0


def test_runner_unique_runs_and_log_artifacts() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_prod_test_"))
    if os.name == "nt":
        exe = tmp / "mock_iitpave.cmd"
        exe.write_text(
            "@echo off\n"
            "echo IITPAVE OUTPUT\n"
            "echo Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR\n"
            "echo 200 0 0.35 0.11 0.09 0 0 -0.000050 0.000120 0.000110\n"
            "echo 680L 0 0.04 0.01 0.01 0 0 -0.000220 -0.000030 -0.000020\n",
            encoding="utf-8",
        )
    else:
        exe = tmp / "mock_iitpave"
        exe.write_text(
            "#!/bin/sh\n"
            f"echo '{IITPAVE_OUTPUT_TABLE}'\n",
            encoding="utf-8",
        )
        exe.chmod(0o755)

    runner = IITPaveExternalExeRunner(exe_path=exe, timeout_sec=10.0, use_stdin_stdout=True)
    out = runner.run("dummy input text")
    assert "IITPAVE OUTPUT" in out

    # Check that runs folder has the run directory
    runs_dir = USER_DATA_DIR / "iitpave_runs"
    assert runs_dir.is_dir()
    
    # Find the latest run folder
    subdirs = sorted(runs_dir.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True)
    assert len(subdirs) > 0
    latest_run = subdirs[0]

    # Verify all 5 files are stored
    assert (latest_run / "iitp_inp.dat").is_file()
    assert (latest_run / "iitp_out.dat").is_file()
    assert (latest_run / "stdout.log").is_file()
    assert (latest_run / "stderr.log").is_file()
    assert (latest_run / "run_status.json").is_file()

    # Verify run status fields
    status = json.loads((latest_run / "run_status.json").read_text(encoding="utf-8"))
    assert status["status"] == "success"
    assert status["returncode"] == 0
    assert status["duration_sec"] > 0


def test_stabilized_workflow_mechanistic_verified_mode() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_prod_test_"))
    if os.name == "nt":
        exe = tmp / "mock_iitpave.cmd"
        exe.write_text(
            "@echo off\n"
            "echo IITPAVE OUTPUT\n"
            "echo Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR\n"
            "echo 200 0 0.35 0.11 0.09 0 0 -0.000050 0.000120 0.000110\n"
            "echo 680L 0 0.04 0.01 0.01 0 0 -0.000220 -0.000030 -0.000020\n",
            encoding="utf-8",
        )
    else:
        exe = tmp / "mock_iitpave"
        exe.write_text(
            "#!/bin/sh\n"
            f"echo '{IITPAVE_OUTPUT_TABLE}'\n",
            encoding="utf-8",
        )
        exe.chmod(0o755)

    inp = StabilizedInput(
        ctb_thickness_mm=120.0,
        ctb_modulus_mpa=5000.0,
        ctb_ucs_mpa=4.5,
        ctb_poisson=0.25,
        cts_class="C1.5/2.0",
        cts_thickness_mm=100.0,
        cts_modulus_mpa=3000.0,
        cts_poisson=0.25,
        bituminous_thickness_mm=100.0,
        gsb_thickness_mm=150.0,
        flexible_design_msa=10.0,
        flexible_subgrade_cbr=5.0,
        notes="Test run",
    )
    res = compute_stabilized_design(inp, has_mechanistic_validation=False)

    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_EXTERNAL,
        configured_executable_path=str(exe),
    )
    workflow = run_stabilized_iitpave_mechanistic_workflow(res, runner_config=cfg)
    
    assert workflow.summary is not None
    assert workflow.summary.refused is False
    assert workflow.structural_result.validation_mode == "Mechanistic Verified Mode"
    assert workflow.structural_result.mechanistic_validation is not None


def test_stabilized_workflow_decision_support_mode_on_stub() -> None:
    inp = StabilizedInput(
        ctb_thickness_mm=120.0,
        ctb_modulus_mpa=5000.0,
        ctb_ucs_mpa=4.5,
        ctb_poisson=0.25,
        cts_class="C1.5/2.0",
        cts_thickness_mm=100.0,
        cts_modulus_mpa=3000.0,
        cts_poisson=0.25,
        bituminous_thickness_mm=100.0,
        gsb_thickness_mm=150.0,
        flexible_design_msa=10.0,
        flexible_subgrade_cbr=5.0,
        notes="Test stub run",
    )
    res = compute_stabilized_design(inp, has_mechanistic_validation=False)

    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_STUB,
    )
    workflow = run_stabilized_iitpave_mechanistic_workflow(res, runner_config=cfg)

    # Stub run returns summary marked is_placeholder=True, which keeps it in Decision Support Mode
    assert workflow.summary is not None
    assert workflow.summary.is_placeholder is True
    assert workflow.structural_result.validation_mode == "Decision Support Mode"


def test_stabilized_workflow_parser_failure_raises_exact_string() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_prod_test_"))
    if os.name == "nt":
        exe = tmp / "mock_iitpave.cmd"
        exe.write_text(
            "@echo off\n"
            "echo IITPAVE COMPLETED BUT UNPARSABLE GARBAGE OUTPUT\n",
            encoding="utf-8",
        )
    else:
        exe = tmp / "mock_iitpave"
        exe.write_text(
            "#!/bin/sh\n"
            "echo 'IITPAVE COMPLETED BUT UNPARSABLE GARBAGE OUTPUT'\n",
            encoding="utf-8",
        )
        exe.chmod(0o755)

    inp = StabilizedInput(
        ctb_thickness_mm=120.0,
        ctb_modulus_mpa=5000.0,
        ctb_ucs_mpa=4.5,
        ctb_poisson=0.25,
        cts_class="C1.5/2.0",
        cts_thickness_mm=100.0,
        cts_modulus_mpa=3000.0,
        cts_poisson=0.25,
        bituminous_thickness_mm=100.0,
        gsb_thickness_mm=150.0,
        flexible_design_msa=10.0,
        flexible_subgrade_cbr=5.0,
        notes="Test parser failure",
    )
    res = compute_stabilized_design(inp, has_mechanistic_validation=False)

    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_EXTERNAL,
        configured_executable_path=str(exe),
    )
    workflow = run_stabilized_iitpave_mechanistic_workflow(res, runner_config=cfg)

    # Execution completed but output contract/parser checks must block it
    assert workflow.blocked is True
    assert workflow.blocked_reason == "IITPAVE execution completed but output verification failed."
    assert workflow.structural_result.validation_mode == "Decision Support Mode"


def test_iitpave_executable_path_validation() -> None:
    # Test discovery path validation
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_test_val_"))
    dummy_exe = tmp / "IITPAVE.exe"
    dummy_exe.write_text("Dummy binary content", encoding="utf-8")
    
    validation = validate_iitpave_environment(configured_path=str(dummy_exe))
    assert validation.selected_path == dummy_exe
    assert validation.ok is True


def test_iitpave_output_parser_cases() -> None:
    # Test parsing real output table
    output_text = (
        "IITPAVE REGION RESULTS\n"
        " Z      R    SigmaZ   SigmaT   SigmaR   epZ      epT      epR\n"
        " 100    0    0.45     0.15     0.10     -50      110      90\n"
        " 250    0    0.15     0.05     0.04     -200     20       10\n"
    )
    from app.core.iitpave.parser import parse_iitpave_output
    from app.core.iitpave.runner import SOURCE_EXTERNAL
    
    result = parse_iitpave_output(output_text, source=SOURCE_EXTERNAL)
    assert len(result.point_results) == 2
    assert result.point_results[0].z_mm == 100.0
    assert result.point_results[0].epsilon_t_microstrain == 110.0
    assert result.point_results[1].epsilon_z_microstrain == -200.0


def test_iitpave_optimization_loop_fatigue_governs() -> None:
    # Test iterative optimization under fatigue failure governs
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_opt_fatigue_"))
    
    # We write a mock executable that checks layer 2 (DBM) thickness.
    # If DBM < 100mm, fatigue life remains unsafe.
    # Once DBM >= 100mm, fatigue life passes.
    python_cmd = (
        "import sys\n"
        "lines = [ln.strip() for ln in open('iitp_inp.dat').read().splitlines() if ln.strip() and not ln.strip().startswith('#')]\n"
        "n_layers = int(lines[0])\n"
        "dbm_thick = 0.0\n"
        "for i in range(1, n_layers + 1):\n"
        "    t = lines[i].split()\n"
        "    if len(t) >= 3 and i == 2: dbm_thick = float(t[2])\n"
        "if dbm_thick < 100.0:\n"
        "    print('IITPAVE OUTPUT')\n"
        "    print('Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR')\n"
        "    print('40 0 0.35 0.11 0.09 0 0 -0.000850 0.000850 0.000110')\n"
        "    print('300 0 0.04 0.01 0.01 0 0 -0.001220 -0.000030 -0.000020')\n"
        "else:\n"
        "    print('IITPAVE OUTPUT')\n"
        "    print('Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR')\n"
        "    print('40 0 0.35 0.11 0.09 0 0 -0.000050 0.000040 0.000030')\n"
        "    print('300 0 0.04 0.01 0.01 0 0 -0.000020 -0.000030 -0.000020')\n"
    )
    
    py_file = tmp / "mock_opt.py"
    py_file.write_text(python_cmd, encoding="utf-8")
    
    if os.name == "nt":
        exe = tmp / "mock_opt.cmd"
        exe.write_text(f"@echo off\npython \"{py_file}\"\n", encoding="utf-8")
    else:
        exe = tmp / "mock_opt"
        exe.write_text(f"#!/bin/sh\npython \"{py_file}\"\n", encoding="utf-8")
        exe.chmod(0o755)
        
    from app.core import StructuralInput, StructuralResult, PavementLayer
    from app.db.repository import Database
    
    layers = (
        PavementLayer(name="BC", thickness_mm=40.0, material="BC", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="DBM", thickness_mm=80.0, material="DBM", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="WMM", thickness_mm=250.0, material="WMM", modulus_mpa=300.0, poisson=0.4),
        PavementLayer(name="GSB", thickness_mm=150.0, material="GSB", modulus_mpa=150.0, poisson=0.4)
    )
    inputs = StructuralInput(
        road_category="NH / SH", design_life_years=15, initial_cvpd=2000.0,
        growth_rate_pct=7.5, vdf=2.5, ldf=0.75, subgrade_cbr_pct=5.0
    )
    result = StructuralResult(
        inputs=inputs, design_msa=10.0, growth_factor=18.0, subgrade_mr_mpa=50.0,
        composition=layers, total_pavement_thickness_mm=520.0
    )
    
    db_file = tmp / "test_opt.db"
    db = Database(db_file)
    
    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_EXTERNAL,
        configured_executable_path=str(exe),
        bc_min=30.0, bc_max=80.0, dbm_min=50.0, dbm_max=300.0,
        wmm_max=250.0, gsb_max=400.0, max_iterations=10
    )
    
    from app.core.iitpave.workflow import run_structural_iitpave_optimization_workflow
    wf_res, iterations = run_structural_iitpave_optimization_workflow(
        result, db=db, project_id=1, runner_config=cfg
    )
    
    # Strains should pass in step 2 (DBM increased from 80mm to 90mm then 100mm)
    assert len(iterations) >= 2
    assert wf_res.ok is True
    final_dbm = next(ly.thickness_mm for ly in wf_res.structural_result.composition if ly.name.upper() == "DBM")
    assert final_dbm >= 100.0


def test_iitpave_optimization_loop_rutting_governs() -> None:
    # Test iterative optimization under rutting failure governs
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_opt_rutting_"))
    
    # Mock script: if GSB < 200mm, rutting remains unsafe.
    python_cmd = (
        "import sys\n"
        "lines = [ln.strip() for ln in open('iitp_inp.dat').read().splitlines() if ln.strip() and not ln.strip().startswith('#')]\n"
        "n_layers = int(lines[0])\n"
        "gsb_thick = 0.0\n"
        "for i in range(1, n_layers + 1):\n"
        "    t = lines[i].split()\n"
        "    if len(t) >= 3 and i == 4: gsb_thick = float(t[2])\n"
        "if gsb_thick < 200.0:\n"
        "    print('IITPAVE OUTPUT')\n"
        "    print('Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR')\n"
        "    print('40 0 0.35 0.11 0.09 0 0 -0.000050 0.000040 0.000030')\n"
        "    print('300 0 0.04 0.01 0.01 0 0 -0.000850 -0.000030 -0.000020')\n"
        "else:\n"
        "    print('IITPAVE OUTPUT')\n"
        "    print('Z R SigmaZ SigmaT SigmaR TaoRZ DispZ epZ epT epR')\n"
        "    print('40 0 0.35 0.11 0.09 0 0 -0.000050 0.000040 0.000030')\n"
        "    print('300 0 0.04 0.01 0.01 0 0 -0.000080 -0.000030 -0.000020')\n"
    )
    
    py_file = tmp / "mock_opt_rut.py"
    py_file.write_text(python_cmd, encoding="utf-8")
    
    if os.name == "nt":
        exe = tmp / "mock_opt_rut.cmd"
        exe.write_text(f"@echo off\npython \"{py_file}\"\n", encoding="utf-8")
    else:
        exe = tmp / "mock_opt_rut"
        exe.write_text(f"#!/bin/sh\npython \"{py_file}\"\n", encoding="utf-8")
        exe.chmod(0o755)
        
    from app.core import StructuralInput, StructuralResult, PavementLayer
    from app.db.repository import Database
    
    layers = (
        PavementLayer(name="BC", thickness_mm=40.0, material="BC", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="DBM", thickness_mm=80.0, material="DBM", modulus_mpa=3000.0, poisson=0.35),
        # Set WMM to max 250mm so GSB is increased next
        PavementLayer(name="WMM", thickness_mm=250.0, material="WMM", modulus_mpa=300.0, poisson=0.4),
        PavementLayer(name="GSB", thickness_mm=150.0, material="GSB", modulus_mpa=150.0, poisson=0.4)
    )
    inputs = StructuralInput(
        road_category="NH / SH", design_life_years=15, initial_cvpd=2000.0,
        growth_rate_pct=7.5, vdf=2.5, ldf=0.75, subgrade_cbr_pct=5.0
    )
    result = StructuralResult(
        inputs=inputs, design_msa=10.0, growth_factor=18.0, subgrade_mr_mpa=50.0,
        composition=layers, total_pavement_thickness_mm=520.0
    )
    
    db_file = tmp / "test_opt_rut.db"
    db = Database(db_file)
    
    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_EXTERNAL,
        configured_executable_path=str(exe),
        bc_min=30.0, bc_max=80.0, dbm_min=50.0, dbm_max=300.0,
        wmm_max=250.0, gsb_max=400.0, max_iterations=10
    )
    
    from app.core.iitpave.workflow import run_structural_iitpave_optimization_workflow
    wf_res, iterations = run_structural_iitpave_optimization_workflow(
        result, db=db, project_id=1, runner_config=cfg
    )
    
    # Rutting governs, WMM is at max (250mm), so GSB is adjusted.
    assert len(iterations) >= 2
    assert wf_res.ok is True
    final_gsb = next(ly.thickness_mm for ly in wf_res.structural_result.composition if ly.name.upper() == "GSB")
    assert final_gsb >= 200.0


def test_database_persistence_run_logs() -> None:
    # Test persistence of executable path, input/output paths, stdout, and iteration details
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_db_test_"))
    db_file = tmp / "test_persistence.db"
    
    from app.db.repository import Database
    from app.db.schema import Project
    from app.core.mechanistic_validation import MechanisticValidationSummary, FatigueCheck, RuttingCheck, FatigueCalibration, RuttingCalibration
    
    db = Database(db_file)
    with db.session() as s:
        p = Project(work_name="Traceability Project")
        s.add(p)
        s.flush()
        pid = p.id
        
    f_cal = FatigueCalibration(label="Test", k1=0.05, k2=0.05, k3=0.05, reliability_pct=80, is_placeholder=False)
    r_cal = RuttingCalibration(label="Test", k_r=0.05, k_v=0.05, reliability_pct=80, is_placeholder=False)
    
    fatigue = FatigueCheck("PASS", 12.0, 10.0, 45.0, False, "", f_cal)
    rutting = RuttingCheck("PASS", 15.0, 10.0, 120.0, False, "", r_cal)
    summary = MechanisticValidationSummary(fatigue, rutting, is_placeholder=False, refused=False, refused_reason="")
    
    iterations = [
        {"attempt": 1, "verdict": "FAIL", "decision": "Increase DBM"},
        {"attempt": 2, "verdict": "PASS", "decision": "Keep"}
    ]
    recommended_thick = {"BC": 40, "DBM": 100}
    
    row = db.save_mechanistic_validation(
        project_id=pid,
        summary=summary,
        inputs={"dummy": True},
        exe_path="C:/IITPAVE/IITPAVE.exe",
        input_filepath="C:/IITPAVE/iitp_inp.dat",
        output_filepath="C:/IITPAVE/iitp_out.dat",
        stdout_log="Calculation success console output",
        stderr_log="No errors logged",
        iteration_history_json=json.dumps(iterations),
        recommended_thickness_json=json.dumps(recommended_thick)
    )
    
    # Reload and assert
    loaded = db.latest_mechanistic_validation(pid)
    assert loaded is not None
    assert loaded.exe_path == "C:/IITPAVE/IITPAVE.exe"
    assert loaded.input_filepath == "C:/IITPAVE/iitp_inp.dat"
    assert loaded.output_filepath == "C:/IITPAVE/iitp_out.dat"
    assert loaded.stdout_log == "Calculation success console output"
    assert loaded.stderr_log == "No errors logged"
    assert "attempt" in loaded.iteration_history_json
    assert "BC" in loaded.recommended_thickness_json


def test_missing_iitpave_does_not_crash_and_saves_refused_cleanly() -> None:
    # Test that missing IITPAVE does not crash the workflow and saves refused cleanly
    from app.core.iitpave.workflow import run_structural_iitpave_mechanistic_workflow
    from app.core import StructuralInput, StructuralResult, PavementLayer
    from app.db.repository import Database
    import tempfile
    
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_test_missing_"))
    db_file = tmp / "test_missing.db"
    db = Database(db_file)
    
    layers = (
        PavementLayer(name="BC", thickness_mm=40.0, material="BC", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="DBM", thickness_mm=80.0, material="DBM", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="WMM", thickness_mm=250.0, material="WMM", modulus_mpa=300.0, poisson=0.4),
        PavementLayer(name="GSB", thickness_mm=150.0, material="GSB", modulus_mpa=150.0, poisson=0.4)
    )
    inputs = StructuralInput(
        road_category="NH / SH", design_life_years=15, initial_cvpd=2000.0,
        growth_rate_pct=7.5, vdf=2.5, ldf=0.75, subgrade_cbr_pct=5.0
    )
    result = StructuralResult(
        inputs=inputs, design_msa=10.0, growth_factor=18.0, subgrade_mr_mpa=50.0,
        composition=layers, total_pavement_thickness_mm=520.0
    )
    
    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_EXTERNAL,
        configured_executable_path="C:/non_existent_iitpave_path_12345/IITPAVE.exe",
    )
    
    # 1. Verification run should return a blocked/refused result gracefully without crashing
    workflow_res = run_structural_iitpave_mechanistic_workflow(result, runner_config=cfg)
    assert workflow_res.blocked is True
    assert "not found" in workflow_res.blocked_reason.lower() or "not exist" in workflow_res.blocked_reason.lower() or "no usable local" in workflow_res.blocked_reason.lower() or "invalid" in workflow_res.blocked_reason.lower()
    
    # 2. summary is None when blocked
    assert workflow_res.summary is None
    
    # 3. Database persistence should work cleanly
    row = db.save_mechanistic_validation(
        project_id=1,
        summary=workflow_res.summary,
        inputs=workflow_res.as_dict()
    )
    assert row is not None
    assert row.refused is True
    assert row.fatigue_verdict is None
    assert row.rutting_verdict is None


def test_stub_missing_iitpave_never_produces_pass_verdict() -> None:
    from app.core.iitpave.workflow import run_structural_iitpave_mechanistic_workflow
    from app.core import StructuralInput, StructuralResult, PavementLayer
    
    layers = (
        PavementLayer(name="BC", thickness_mm=40.0, material="BC", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="DBM", thickness_mm=80.0, material="DBM", modulus_mpa=3000.0, poisson=0.35),
        PavementLayer(name="WMM", thickness_mm=250.0, material="WMM", modulus_mpa=300.0, poisson=0.4),
        PavementLayer(name="GSB", thickness_mm=150.0, material="GSB", modulus_mpa=150.0, poisson=0.4)
    )
    inputs = StructuralInput(
        road_category="NH / SH", design_life_years=15, initial_cvpd=2000.0,
        growth_rate_pct=7.5, vdf=2.5, ldf=0.75, subgrade_cbr_pct=5.0
    )
    result = StructuralResult(
        inputs=inputs, design_msa=10.0, growth_factor=18.0, subgrade_mr_mpa=50.0,
        composition=layers, total_pavement_thickness_mm=520.0
    )
    
    cfg = IITPaveRunnerConfig(
        mode=IITPAVE_RUNNER_STUB,
    )
    
    workflow_res = run_structural_iitpave_mechanistic_workflow(result, runner_config=cfg)
    
    # Verdicts MUST be None (which maps to N/A), never PASS or FAIL
    assert workflow_res.summary is not None
    assert workflow_res.summary.refused is True
    assert workflow_res.summary.fatigue.verdict is None
    assert workflow_res.summary.rutting.verdict is None


def test_decision_support_mode_locking_rules() -> None:
    # Test locking checks in submission center panel rehydration logic
    import tempfile
    from app.db.repository import Database
    from app.db.schema import Project
    
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_test_lock_"))
    db_file = tmp / "test_lock.db"
    db = Database(db_file)
    
    # Create project
    proj = db.create_project(work_name="Lock Verification Project")
    db.initialize_workflow_statuses(proj.id)
    
    # Check that initially mechanistic validation is not run
    mech_val = db.latest_mechanistic_validation(proj.id)
    assert mech_val is None
    
    # Lock is allowed if checklist has "IRC Catalogue Design (Decision Support Mode)"
    # We can mock this checklist save
    db.save_project_checklist(proj.id, "Draft", {
        "iitpave_verification": "IRC Catalogue Design (Decision Support Mode)"
    })
    
    db.lock_project(proj.id)
    p = db.get_project(proj.id)
    assert p.locked is True


def test_reports_show_correct_unavailable_wording() -> None:
    import tempfile
    from app.db.repository import Database
    from app.db.schema import Project, StructuralDesign
    from app.reports.report_builder import build_combined_report, CombinedReportContext
    
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_test_report_"))
    db_file = tmp / "test_report.db"
    db = Database(db_file)
    
    proj = db.create_project(work_name="Report Wording Project")
    db.initialize_workflow_statuses(proj.id)
    
    # Creating a stub structural design to bypass build_combined_report raise check
    with db.session() as s:
        sd = StructuralDesign(
            project_id=proj.id,
            inputs_json="{}",
            composition_json="[]",
            total_pavement_thickness_mm=0.0
        )
        s.add(sd)
        s.commit()
        
    meta = {
        "project_title": proj.work_name,
        "work_name": proj.work_name,
        "work_order_no": "",
        "work_order_date": "",
        "client": "",
        "agency": "",
        "submitted_by": "",
        "report_date": "29-Jun-2026",
        "binder_grade": "",
        "mix_type_key": "",
    }
    ctx = CombinedReportContext(**meta)
    out_path = tmp / "report_wording.docx"
    
    build_combined_report(out_path, db, proj.id, ctx)
    assert out_path.is_file()
    
    # Read the text of docx to verify wording exists
    import docx
    doc = docx.Document(out_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
                
    combined_text = "\n".join(full_text)
    assert "Mechanistic verification was not executed because a licensed IITPAVE installation was unavailable." in combined_text


def test_common_folder_discovery() -> None:
    # Test auto-detection from common folders
    from app.core import discover_iitpave_executable
    import tempfile
    from pathlib import Path
    
    tmp = Path(tempfile.mkdtemp(prefix="iitpave_test_common_"))
    
    # Create fake executable in common folders candidates structure
    fake_common = tmp / "IITPAVE.exe"
    fake_common.write_text("fake binary", encoding="utf-8")
    
    # Inject fake_common path as one of the common folders search candidates
    # We can mock common_iitpave_exe_candidates to return [fake_common]
    import app.core.iitpave.discovery as discovery
    old_candidates = discovery.common_iitpave_exe_candidates
    discovery.common_iitpave_exe_candidates = lambda: [fake_common]
    
    try:
        candidates = discover_iitpave_executable(configured_path=None)
        # Find if our mock path was discovered under SOURCE_COMMON
        common_candidates = [c for c in candidates if c.source == "common_folders"]
        assert len(common_candidates) == 1
        assert common_candidates[0].path == fake_common
    finally:
        discovery.common_iitpave_exe_candidates = old_candidates


