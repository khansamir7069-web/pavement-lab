"""Phase-36 smoke - selected IITPAVE schema history report inclusion."""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_tmp = Path(tempfile.mkdtemp(prefix="phase36_iitpave_schema_history_selection_"))
_db_path = _tmp / "phase36.db"

import app.config as _cfg  # noqa: E402
_cfg.DB_PATH = _db_path

from docx import Document  # noqa: E402
from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QFileDialog, QDialog, QListWidget, QMessageBox  # noqa: E402

from app.core import DEFAULT_FIXTURE_MANIFEST, FIXTURE_STATUS_VERIFIED  # noqa: E402
from app.db.project_exchange import export_project, import_project  # noqa: E402
from app.db.repository import Database  # noqa: E402
from app.reports import (  # noqa: E402
    CombinedReportContext,
    IITPAVE_SCHEMA_HISTORY_SELECTION_STATUS_PARTIAL_UNKNOWN,
    build_combined_report,
    build_iitpave_schema_history_inclusion_selection,
    build_iitpave_schema_history_report_summary,
    build_iitpave_schema_history_review,
    run_iitpave_schema_diagnostics_workflow,
)

import app.db.repository as _repo  # noqa: E402


_messages: list[tuple[str, str, str]] = []
QMessageBox.information = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("info", title, message)) or 0
)
QMessageBox.warning = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("warning", title, message)) or 0
)
QMessageBox.critical = staticmethod(
    lambda _parent, title, message, *a, **k: _messages.append(("critical", title, message)) or 0
)


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _fixture_folder(name: str) -> Path:
    folder = _tmp / name
    folder.mkdir(exist_ok=True)
    sample = _write_fixture(folder, f"{name}.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(folder, [{
        "filename": sample.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    return folder


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _persist_history_pair(db: Database, project_id: int) -> tuple:
    first = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("first_fixture"),
        report_path=_tmp / "first_schema_diagnostics.docx",
    )
    second = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("second_fixture"),
        report_path=_tmp / "second_schema_diagnostics.docx",
    )
    row1 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=first)
    row2 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=second)
    return row1, row2


def _safe_text(value) -> str:
    return json.dumps(value, sort_keys=True).lower()


def main() -> int:
    db = Database(_db_path)
    project = db.create_project(
        work_name="Phase 36 History Selection",
        agency="State PWD",
        submitted_by="QA Engineer",
        mix_type="DBM-II",
    )
    row1, row2 = _persist_history_pair(db, project.id)
    rows = db.list_iitpave_schema_diagnostics(project.id)

    print("=== 1) Typed selection preserves deterministic history order ===")
    selection = build_iitpave_schema_history_inclusion_selection(
        project.id,
        rows,
        selected_history_ids=(999999, row1.id),
    )
    assert selection.status == IITPAVE_SCHEMA_HISTORY_SELECTION_STATUS_PARTIAL_UNKNOWN
    assert selection.available_history_ids == (row2.id, row1.id)
    assert selection.selected_history_ids == (row1.id,)
    assert selection.skipped_unknown_history_ids == (999999,)
    selection_text = _safe_text(selection.as_dict())
    assert '"engineering_calculations_allowed": false' in selection_text
    assert "fatigue" not in selection_text
    assert "rutting" not in selection_text
    assert "irc_compliance" not in selection_text
    assert "point_results" not in selection_text
    print("  [PASS] selection model is deterministic and audit-only")

    print("\n=== 2) Report summary reflects selected history only ===")
    review = build_iitpave_schema_history_review(
        project.id,
        rows,
        selected_history_ids=(row1.id,),
    )
    summary = build_iitpave_schema_history_report_summary(review)
    assert review.item_count == 1
    assert review.latest_item is not None
    assert review.latest_item.id == row1.id
    assert summary.included_history_count == 1
    assert summary.engineering_calculations_allowed is False
    assert summary.selection is not None
    assert summary.selection.selected_history_ids == (row1.id,)
    print("  [PASS] typed review/report summary follows operator selection")

    print("\n=== 3) Combined report includes only selected history records ===")
    selected_doc = _tmp / "selected_schema_history.docx"
    _out, included = build_combined_report(
        selected_doc,
        db,
        project.id,
        CombinedReportContext(
            project_title=project.work_name,
            work_name=project.work_name,
            agency=project.agency or "",
            submitted_by=project.submitted_by or "",
            mix_type_key=project.mix_type or "",
        ),
        schema_history_selection_ids=(row1.id,),
    )
    assert "IITPAVE Schema Diagnostics History" in included
    text = _doc_text(selected_doc)
    assert f"#{row1.id}" in text
    assert f"#{row2.id}" not in text
    assert "Report Inclusion Selection" in text
    assert "Selected records" in text
    assert "Engineering calculations allowed" in text
    assert "fatigue_life" not in text.lower()
    assert "rutting_life" not in text.lower()
    assert "irc_compliance" not in text.lower()
    assert "point_results" not in text.lower()
    print("  [PASS] report output respects selected history IDs")

    print("\n=== 4) Empty selection blocks history-only report without calculations ===")
    try:
        build_combined_report(
            _tmp / "empty_selection.docx",
            db,
            project.id,
            CombinedReportContext(work_name=project.work_name, mix_type_key=project.mix_type or ""),
            schema_history_selection_ids=(),
        )
    except ValueError as exc:
        assert "No module data found" in str(exc)
    else:
        raise AssertionError("Empty schema history selection should not create a report.")
    print("  [PASS] empty selection does not fabricate a report section")

    print("\n=== 5) Export/import compatibility keeps selected history reportable ===")
    exported = export_project(db, project.id)
    imported = import_project(db, exported)
    imported_rows = db.list_iitpave_schema_diagnostics(imported.project_id)
    imported_selected = imported_rows[-1]
    imported_doc = _tmp / "imported_selected_schema_history.docx"
    _out, imported_included = build_combined_report(
        imported_doc,
        db,
        imported.project_id,
        CombinedReportContext(work_name=imported.work_name, mix_type_key="DBM-II"),
        schema_history_selection_ids=(imported_selected.id,),
    )
    assert "IITPAVE Schema Diagnostics History" in imported_included
    assert f"#{imported_selected.id}" in _doc_text(imported_doc)
    print("  [PASS] imported records remain selectable for report inclusion")

    print("\n=== 6) UI combined-report workflow exposes checkable history selection ===")
    _repo._singleton = db
    app = QApplication.instance() or QApplication([])
    from app.ui.main_window import MainWindow  # noqa: E402

    original_exec = QDialog.exec

    def _accept_only_first(dialog, *args, **kwargs):
        selector = dialog.findChild(QListWidget)
        if selector is not None:
            for idx in range(selector.count()):
                item = selector.item(idx)
                item.setCheckState(Qt.Checked if idx == 0 else Qt.Unchecked)
        return QDialog.Accepted

    QDialog.exec = _accept_only_first
    w = MainWindow()
    w._current_project_id = project.id
    ui_report = _tmp / "ui_selected_schema_history.docx"
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(ui_report), "Word"))
    w._on_export_combined_report(project.id)
    QDialog.exec = original_exec
    ui_text = _doc_text(ui_report)
    assert f"#{row2.id}" in ui_text
    assert f"#{row1.id}" not in ui_text
    assert any(item[1] == "Combined report exported" for item in _messages)
    w.close()
    app.processEvents()
    print("  [PASS] UI selection controls feed the combined report builder")

    print("\nPHASE 36 IITPAVE SCHEMA HISTORY SELECTION SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
