"""Phase-39 smoke - consultancy report provenance and traceability section."""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_tmp = Path(tempfile.mkdtemp(prefix="phase39_report_provenance_"))
_db_path = _tmp / "phase39.db"

import app.config as _cfg  # noqa: E402
_cfg.DB_PATH = _db_path

from docx import Document  # noqa: E402

from app.core import DEFAULT_FIXTURE_MANIFEST, FIXTURE_STATUS_VERIFIED  # noqa: E402
from app.db.project_exchange import export_project, import_project  # noqa: E402
from app.db.repository import Database  # noqa: E402
from app.reports import (  # noqa: E402
    CombinedReportContext,
    CombinedReportProvenanceSummary,
    build_combined_report,
    run_iitpave_schema_diagnostics_workflow,
)


def _write_manifest(folder: Path, fixtures: list[dict]) -> None:
    (folder / DEFAULT_FIXTURE_MANIFEST).write_text(
        json.dumps({"fixtures": fixtures}, indent=2),
        encoding="utf-8",
    )


def _write_fixture(folder: Path, name: str, lines: list[str]) -> Path:
    path = folder / name
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _fixture_folder(name: str) -> Path:
    folder = _tmp / name
    folder.mkdir(exist_ok=True)
    sample = _write_fixture(folder, f"{name}.out", [
        "IITPAVE OUTPUT",
        "Stress and Strain table",
        "Layer Radius Value",
        "1 0.0 12.3",
    ])
    _write_manifest(folder, [{
        "filename": sample.name,
        "verification_status": FIXTURE_STATUS_VERIFIED,
    }])
    return folder


def _doc_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _persist_history_pair(db: Database, project_id: int) -> tuple:
    first = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("first_fixture"),
        report_path=_tmp / "first_schema_diagnostics.docx",
    )
    second = run_iitpave_schema_diagnostics_workflow(
        _fixture_folder("second_fixture"),
        report_path=_tmp / "second_schema_diagnostics.docx",
    )
    row1 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=first)
    row2 = db.save_iitpave_schema_diagnostics(project_id=project_id, result=second)
    return row1, row2


def _safe_text(value) -> str:
    return json.dumps(value, sort_keys=True).lower()


def main() -> int:
    db = Database(_db_path)
    project = db.create_project(
        work_name="Phase 39 Provenance Report",
        agency="State PWD",
        submitted_by="QA Engineer",
        mix_type="DBM-II",
    )
    row1, row2 = _persist_history_pair(db, project.id)

    print("=== 1) Provenance section renders safely with warning metadata ===")
    report = _tmp / "phase39_provenance.docx"
    _out, included = build_combined_report(
        report,
        db,
        project.id,
        CombinedReportContext(
            project_title=project.work_name,
            work_name=project.work_name,
            agency=project.agency or "",
            submitted_by=project.submitted_by or "",
            mix_type_key=project.mix_type or "",
        ),
        schema_history_selection_ids=(999999, row1.id),
    )
    assert included == ["IITPAVE Schema Diagnostics History"]
    text = _doc_text(report)
    assert "REPORT PROVENANCE AND TRACEABILITY" in text
    assert "Generation Metadata" in text
    assert "Report generated at" in text
    assert "Operator identifier" in text
    assert "QA Engineer" in text
    assert "IITPAVE Schema-History Selection Summary" in text
    assert f"{row1.id}" in text
    assert f"{row2.id}" in text
    assert "Unknown requested IDs" in text
    assert "999999" in text
    assert "Validation Warning Summary" in text
    assert "unknown requested history IDs were ignored" in text
    assert "Export And Import Provenance" in text
    assert "source record IDs are preserved only as provenance" in text
    assert "Report Environment Metadata" in text
    assert "Runtime: Python" in text
    assert "fatigue_life" not in text.lower()
    assert "rutting_life" not in text.lower()
    assert "irc_compliance" not in text.lower()
    assert "point_results" not in text.lower()
    print("  [PASS] report provenance renders with selection and warning metadata")

    print("\n=== 2) Missing provenance metadata falls back safely ===")
    sparse_project = db.create_project(work_name="Phase 39 Sparse Project", mix_type="DBM-II")
    sparse_row, _ = _persist_history_pair(db, sparse_project.id)
    sparse_report = _tmp / "phase39_sparse_provenance.docx"
    build_combined_report(
        sparse_report,
        db,
        sparse_project.id,
        CombinedReportContext(mix_type_key=sparse_project.mix_type or ""),
        schema_history_selection_ids=(sparse_row.id,),
    )
    sparse_text = _doc_text(sparse_report)
    assert "REPORT PROVENANCE AND TRACEABILITY" in sparse_text
    assert "Operator identifier" in sparse_text
    assert "Not recorded" in sparse_text
    assert "No report-time validation warnings recorded." in sparse_text
    print("  [PASS] incomplete metadata is rendered with safe fallbacks")

    print("\n=== 3) Imported-history provenance note is preserved ===")
    exported = export_project(db, project.id)
    imported = import_project(db, exported)
    imported_project = db.get_project(imported.project_id)
    imported_rows = db.list_iitpave_schema_diagnostics(imported.project_id)
    imported_report = _tmp / "phase39_imported_provenance.docx"
    build_combined_report(
        imported_report,
        db,
        imported.project_id,
        CombinedReportContext(
            project_title=imported.work_name,
            work_name=imported.work_name,
            agency=imported_project.agency if imported_project else "",
            submitted_by=imported_project.submitted_by if imported_project else "",
            mix_type_key=imported_project.mix_type if imported_project else "",
        ),
        schema_history_selection_ids=(imported_rows[-1].id,),
    )
    imported_text = _doc_text(imported_report)
    assert "Imported IITPAVE schema-history audit IDs are shown exactly as recorded" in imported_text
    assert "not remapped" in imported_text
    assert "Project exchange format" in imported_text
    print("  [PASS] imported-history provenance is explicit and reportable")

    print("\n=== 4) Provenance serialization remains audit-only ===")
    audit_rows = db.list_iitpave_schema_history_selection_audits(project.id)
    assert audit_rows
    payload = json.loads(audit_rows[0].summary_json)
    serialized = _safe_text(payload)
    assert "engineering_calculations_allowed" in serialized
    assert "fatigue_life" not in serialized
    assert "rutting_life" not in serialized
    assert "irc_compliance" not in serialized
    assert "point_results" not in serialized
    assert CombinedReportProvenanceSummary is not None
    print("  [PASS] existing report/export audit persistence remains compatible")

    print("\nPHASE 39 REPORT PROVENANCE SMOKE: ALL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
