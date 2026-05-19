from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PySide6.QtWidgets import QApplication, QTableWidgetItem

from app.core import compute_material_calc, compute_mix_design
from app.core.models import MixDesignInput, ProjectInfo
from app.graphs import build_chart_set
from app.reports.word_report import ReportContext, build_mix_design_docx
from app.ui.widgets.inputs_panel import GradationTab, InputsPanel


def _app() -> QApplication:
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


def _set_blend(tab: GradationTab, values: dict[str, float]) -> None:
    for name, spin in tab.blend_spins.items():
        value = values.get(name)
        spin.setText("" if value is None else f"{value:.3f}")


def _fill_gradation(tab: GradationTab, name: str, value: float = 50.0) -> None:
    ci = tab._available_aggs.index(name) + 1
    for row in range(tab.table.rowCount()):
        tab.table.setItem(row, ci, QTableWidgetItem(f"{value:g}"))


def _fill_10mm_sg(panel: InputsPanel) -> None:
    for row, value in enumerate((3648, 2363, 2005, 1952)):
        panel.tab_spgr.coarse_10.setItem(row, 0, QTableWidgetItem(str(value)))


def _fill_sg_table(table, values: tuple[float, float, float, float]) -> None:
    for row, value in enumerate(values):
        table.setItem(row, 0, QTableWidgetItem(str(value)))


def _fill_required_lab_inputs(panel: InputsPanel) -> None:
    _fill_sg_table(panel.tab_spgr.coarse_25, (3647, 2363, 2003, 1954))
    _fill_sg_table(panel.tab_spgr.coarse_20, (3649, 2363, 2010, 1951))
    _fill_sg_table(panel.tab_spgr.coarse_10, (3648, 2363, 2005, 1952))
    _fill_sg_table(panel.tab_spgr.fine_6, (492, 834, 1655, 1435))
    _fill_sg_table(panel.tab_spgr.fine_sd, (486, 839, 1664, 1439))
    _fill_sg_table(panel.tab_spgr.bitumen, (40.76, 97.5, 65.56, 97.77))

    gmb_rows = [
        (3.5, 1233, 732, 1236), (3.5, 1231, 730, 1234), (3.5, 1232, 729, 1234),
        (4.0, 1239, 740, 1245), (4.0, 1246, 741, 1249), (4.0, 1247, 741, 1250),
        (4.5, 1241, 741, 1245), (4.5, 1247, 742, 1249), (4.5, 1248, 743, 1251),
        (5.0, 1241, 736, 1243), (5.0, 1239, 735, 1242), (5.0, 1241, 736, 1243),
        (5.5, 1241, 736, 1246), (5.5, 1245, 736, 1247), (5.5, 1243, 735, 1245),
    ]
    for row, (pb, a, c, b) in enumerate(gmb_rows):
        for col, value in ((0, pb), (1, f"S-{row + 1}"), (2, a), (3, c), (4, b)):
            panel.tab_gmb.table.setItem(row, col, QTableWidgetItem(str(value)))

    panel.tab_gmm.pb_ref.setText("4.5")
    for row, values in enumerate(((909, 2152, 3316, 4072), (910, 2154, 3313, 4068))):
        panel.tab_gmm.table.setItem(row, 0, QTableWidgetItem(f"S-{row + 1}"))
        for col, value in enumerate(values, start=1):
            panel.tab_gmm.table.setItem(row, col, QTableWidgetItem(str(value)))

    sf_rows = [
        (3.5, 64, 64, 63.6, 101.9, 1.00, 13.42, 3.21, 12.30),
        (3.5, 64.05, 64, 65, 101.5, 1.00, 13.69, 3.32, None),
        (3.5, 65.71, 65.8, 66.83, 101.73, 0.93, 13.53, 3.55, None),
        (4.0, 65.04, 65.03, 65.09, 101.78, 0.96, 14.28, 3.41, None),
        (4.0, 64.7, 64.32, 64.6, 101.3, 1.00, 14.16, 3.52, None),
        (4.0, 64.72, 64.88, 64.52, 101.82, 0.96, 14.82, 3.38, None),
        (4.5, 62.82, 62.84, 62.4, 101.84, 1.00, 16.33, 3.61, None),
        (4.5, 64.33, 64.54, 64.02, 101.83, 0.96, 16.52, 3.66, None),
        (4.5, 63.01, 63.44, 63.3, 101.73, 1.00, 16.58, 3.59, None),
        (5.0, 64.1, 64.09, 64.2, 101.3, 1.00, 14.31, 3.78, None),
        (5.0, 62.2, 62.29, 62.15, 101.4, 1.04, 14.22, 3.69, None),
        (5.0, 62.35, 62.15, 63, 101.2, 1.04, 13.65, 3.81, None),
        (5.5, 62.15, 62.8, 62.78, 101.6, 1.04, 12.76, 3.86, None),
        (5.5, 62.79, 62.4, 62.7, 101.19, 1.04, 12.7, 3.83, None),
        (5.5, 63.2, 62.7, 62.5, 102.3, 1.00, 14.28, 3.96, None),
    ]
    for row, values in enumerate(sf_rows):
        panel.tab_sf.table.setItem(row, 0, QTableWidgetItem(str(values[0])))
        panel.tab_sf.table.setItem(row, 1, QTableWidgetItem(f"S-{row + 1}"))
        for col, value in enumerate(values[1:], start=2):
            panel.tab_sf.table.setItem(row, col, QTableWidgetItem("" if value is None else str(value)))

    panel.tab_material.std_pb.setText("4.5")
    panel.tab_material.std_agg_wt.setText("1200")
    panel.tab_material.tgt_pb.setText("4.0")


def _docx_text(path) -> str:
    doc = Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_gradation_active_materials_hide_inactive_blend_inputs() -> None:
    _app()
    tab = GradationTab()

    assert tab.material_group.title() == "Active material selection"
    assert tab.material_group.isHidden() is False
    assert tab.active_checks["10mm"].text() == "10 mm"
    assert tab.active_checks["SD"].text() == "Stone Dust"
    assert "Active: 25 mm, 20 mm, 6 mm, Stone Dust, Cement" in tab._active_summary.text()
    assert tab.active_materials() == ("25mm", "20mm", "6mm", "SD", "Cement")
    assert tab.blend_spins["10mm"].isHidden()
    assert tab.table.isColumnHidden(tab._available_aggs.index("10mm") + 1)

    active = ("25mm", "10mm", "6mm", "SD")
    tab.set_active_materials(active)
    _set_blend(tab, {"25mm": 0.30, "10mm": 0.25, "6mm": 0.20, "SD": 0.25})
    for name in ("25mm", "6mm", "SD"):
        _fill_gradation(tab, name)
    _fill_gradation(tab, "10mm")

    grad = tab.collect()

    assert tuple(grad.blend_ratios) == active
    assert tuple(grad.pass_pct) == active
    assert "20mm" not in grad.blend_ratios
    assert "Cement" not in grad.blend_ratios
    assert tab.active_checks["10mm"].isChecked()
    assert "10 mm" in tab._active_summary.text()
    assert "20 mm" not in tab._active_summary.text()
    assert "Cement" not in tab._active_summary.text()
    assert tab.blend_spins["20mm"].isHidden()
    assert tab.blend_spins["Cement"].isHidden()
    assert tab.table.isColumnHidden(tab._available_aggs.index("20mm") + 1)
    assert tab.table.isColumnHidden(tab._available_aggs.index("Cement") + 1)


def test_active_10mm_requires_user_gradation_data() -> None:
    _app()
    tab = GradationTab()
    tab.set_active_materials(("25mm", "10mm"))
    _set_blend(tab, {"25mm": 0.50, "10mm": 0.50})
    _fill_gradation(tab, "25mm")

    try:
        tab.collect()
    except ValueError as exc:
        assert "Missing gradation data for active material 10mm" in str(exc)
    else:
        raise AssertionError("active 10mm without gradation data should be rejected")


def test_mix_inputs_are_blank_by_default_and_validate() -> None:
    _app()
    panel = InputsPanel()

    assert panel.tab_gradation.blend_spins["25mm"].text() == ""
    assert panel.tab_gradation.table.item(0, panel.tab_gradation._available_aggs.index("25mm") + 1).text() == ""
    assert panel.tab_spgr.coarse_25.item(0, 0) is None or panel.tab_spgr.coarse_25.item(0, 0).text() == ""
    assert panel.tab_gmb.table.item(0, 0).text() == ""
    assert panel.tab_gmm.pb_ref.text() == ""
    assert panel.tab_sf.table.item(0, 0).text() == ""
    assert panel.tab_material.std_pb.text() == ""

    panel.tab_gradation.set_active_materials(("25mm", "10mm", "6mm"))
    for tab in (panel.tab_spgr, panel.tab_gmb, panel.tab_gmm, panel.tab_sf, panel.tab_material):
        assert "25 mm, 10 mm, 6 mm" in tab._active_summary.text()

    try:
        panel.collect_all()
    except ValueError as exc:
        assert "Missing blend ratio" in str(exc)
    else:
        raise AssertionError("blank startup inputs should require explicit user data")


def test_inputs_compute_and_report_use_only_active_materials(tmp_path) -> None:
    _app()
    panel = InputsPanel()
    active = ("25mm", "10mm", "6mm", "Cement")
    panel.tab_gradation.set_active_materials(active)
    _set_blend(
        panel.tab_gradation,
        {"25mm": 0.35, "10mm": 0.25, "6mm": 0.20, "Cement": 0.20},
    )
    for name in active:
        _fill_gradation(panel.tab_gradation, name)
    _fill_required_lab_inputs(panel)

    payload = panel.collect_all()
    grad = payload["gradation"]
    coarse, fine, bit = payload["spgr"]

    assert tuple(grad.blend_ratios) == active
    assert set(coarse) == {"25mm", "10mm"}
    assert set(fine) == {"6mm"}
    assert tuple(payload["material_calc"].blend_ratios) == active

    result = compute_mix_design(
        MixDesignInput(
            project=ProjectInfo(mix_type="DBM-II", materials={name: "" for name in active}),
            gradation=grad,
            sg_coarse=coarse,
            sg_fine=fine,
            sg_bitumen=bit,
            gmb=payload["gmb"],
            gmm=payload["gmm"],
            stability_flow=payload["stability_flow"],
        )
    )
    material_calc = compute_material_calc(payload["material_calc"])

    assert tuple(result.gradation.blend_ratios) == active
    assert tuple(row.name for row in material_calc.dry_material_standard) == active

    report_path = tmp_path / "dynamic_materials.docx"
    build_mix_design_docx(
        report_path,
        ReportContext(
            project_title="Dynamic materials",
            mix_type_key="DBM-II",
            materials={name: "" for name in active},
        ),
        result,
        build_chart_set(result.summary, result.obc),
        material_calc=material_calc,
    )
    text = _docx_text(report_path)

    assert "10mm" in text
    assert "20mm" not in text
    assert "SD" not in text
    assert "Cement" in text
